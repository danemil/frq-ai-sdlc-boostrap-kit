#!/usr/bin/env python3
"""The company brand skill (frq-brandbook): its files, its bundled assets, its placement,
and scripts/check_brand.py against a golden off-brand deck and decks built from the
bundled template.

The golden deck (fixtures/frq-brandbook/offbrand-test.pptx) comes from the earlier
internal brand skill's test set: six slides seeded with known violations. The on-brand
decks are built here with zipfile from the template the skill ships, so no python-pptx
is needed (the CI has none); the new_deck.py test runs only where python-pptx is installed.
"""
import ast
import hashlib
import importlib.util
import json
import os
import posixpath
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import helpers
from personal import packs, paths, place

sys.dont_write_bytecode = True     # the tests import the skill's scripts: no __pycache__ in the skill

KIT = helpers.KIT
SKILL = KIT / packs.SKILLS_REL / "frq-brandbook"
CHECK = SKILL / "scripts/check_brand.py"
NEW_DECK = SKILL / "scripts/new_deck.py"
FRQ_PPTX = SKILL / "scripts/frq_pptx.py"
ASSETS = SKILL / "scripts/brand_assets.py"
TEMPLATE = SKILL / "assets/templates/frq-template-slim-core.pptx"
FULL = SKILL / "assets/templates/frq-master.pptx"          # the full 44-layout master (design §8.3)
TOKENS = SKILL / "brand-tokens.json"
GOLDEN = Path(__file__).resolve().parent / "fixtures/frq-brandbook/offbrand-test.pptx"
# The kit owner's skill frq-4-pptx-agent v1.0 as received: every file with its SHA-256, the
# master's parts and the headings of its text (plan 2026-10-09, Task B1 step 1).
INVENTORY = Path(__file__).resolve().parent / "fixtures/frq-brandbook/source-inventory.json"
PLACED = ".agents/skills/ai-sdlc-frq-brandbook"
# Size budgets (design §8.7). The whole folder (it lives once per repo, in .ai-sdlc/kit): 14 MiB.
# The design's "about 13 MB" left out that 0.8.0 already had 2 MB; the merged folder measured
# 14.3 MB (2026-10-10). The placed files (everything .kit-only does not keep back): 0.6 MB.
BUDGET_ALL = 14 * 1024 * 1024
BUDGET_PLACED = 600_000
KIT_ONLY = SKILL / ".kit-only"

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_SLIDE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
REL_LAYOUT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
REL_CHART = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"
CT_SLIDE = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
CT_CHART = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"


def png_alpha(path):
    """(width, height, colour type, [alpha rows]) of an 8-bit RGBA, non-interlaced PNG (stdlib only)."""
    import struct
    import zlib
    b = Path(path).read_bytes()
    i, idat, w, h, ctype = 8, b"", 0, 0, None
    while i < len(b):
        n, typ = struct.unpack(">I4s", b[i:i + 8])
        data = b[i + 8:i + 8 + n]
        i += 12 + n
        if typ == b"IHDR":
            w, h, depth, ctype = struct.unpack(">IIBB", data[:10])
            if (depth, ctype, data[12]) != (8, 6, 0):
                return w, h, (depth, ctype, data[12]), []
        elif typ == b"IDAT":
            idat += data
    raw, bpp = zlib.decompress(idat), 4
    stride, prev, rows = w * bpp, bytearray(w * bpp), []
    for y in range(h):
        f, line = raw[y * (stride + 1)], bytearray(raw[y * (stride + 1) + 1:(y + 1) * (stride + 1)])
        for x in range(stride):
            a = line[x - bpp] if x >= bpp else 0
            up, c = prev[x], (prev[x - bpp] if x >= bpp else 0)
            if f == 1:
                line[x] = (line[x] + a) & 255
            elif f == 2:
                line[x] = (line[x] + up) & 255
            elif f == 3:
                line[x] = (line[x] + (a + up) // 2) & 255
            elif f == 4:
                p_ = a + up - c
                pa, pb, pc = abs(p_ - a), abs(p_ - up), abs(p_ - c)
                line[x] = (line[x] + (a if pa <= pb and pa <= pc else up if pb <= pc else c)) & 255
        rows.append(line[3::4])
        prev = line
    return w, h, (8, 6, 0), rows


def svg_box(path):
    head = Path(path).read_text(encoding="utf-8")[:400]
    vb = [float(x) for x in re.search(r'viewBox="([^"]+)"', head).group(1).split()]
    w, h = (float(re.search(rf'{k}="([\d.]+)"', head).group(1)) for k in ("width", "height"))
    return vb[2] / vb[3], w / h


def run_check(path, *args, python=None):
    """Run check_brand.py as a person would: (exit code, stdout, stderr)."""
    r = subprocess.run([python or sys.executable, "-I", str(CHECK), str(path), *args],
                       capture_output=True, text=True, check=False)
    return r.returncode, r.stdout, r.stderr


def findings(path, *args):
    code, out, err = run_check(path, "--json", *args)
    data = json.loads(out)
    return code, data["findings"]


def layout_file(template, name):
    """ppt/slideLayouts/slideLayoutN.xml of the layout called `name` in `template`."""
    with zipfile.ZipFile(template) as z:
        for n in z.namelist():
            if re.fullmatch(r"ppt/slideLayouts/slideLayout\d+\.xml", n):
                if f'<p:cSld name="{name.replace("&", "&amp;")}"' in z.read(n).decode("utf-8"):
                    return n
    raise KeyError(name)


def sp(ph_type=None, idx=None, text="", rpr="", sppr="", name="Shape"):
    """A shape: a placeholder when ph_type or idx is given, else a text box."""
    ph = ""
    if ph_type or idx is not None:
        attrs = (f' type="{ph_type}"' if ph_type else "") + (f' idx="{idx}"' if idx is not None else "")
        ph = f"<p:nvPr><p:ph{attrs}/></p:nvPr>"
    else:
        ph = "<p:nvPr/>"
    paras = "".join(f"<a:p><a:r>{rpr}<a:t>{line}</a:t></a:r></a:p>" for line in text.split("\n"))
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{abs(hash(name)) % 9000 + 10}" name="{name}"/>'
            f"<p:cNvSpPr/>{ph}</p:nvSpPr><p:spPr>{sppr}</p:spPr>"
            f"<p:txBody><a:bodyPr/><a:lstStyle/>{paras}</p:txBody></p:sp>")


def build_pptx(dest, slides, base=TEMPLATE, footer=True):
    """Copy `base` and append slides: [(layout name, shapes xml, {chart part: xml})].
    `footer`: fill the master footer's title and presenter, as frq_pptx.py and new_deck.py do
    (the template's "Presentation title" and "<by Presenter>" are a check finding)."""
    with zipfile.ZipFile(base) as zin:
        parts = {n: zin.read(n) for n in zin.namelist()}
    if footer:
        master = parts["ppt/slideMasters/slideMaster1.xml"].decode("utf-8")
        master = master.replace("<a:t>Presentation title</a:t>", "<a:t>Remote digital towers</a:t>")
        master = master.replace("<a:t>&lt;by Presenter&gt;</a:t>", "<a:t>by Ana Pop</a:t>")
        parts["ppt/slideMasters/slideMaster1.xml"] = master.encode("utf-8")
    ct = parts["[Content_Types].xml"].decode("utf-8")
    prels = parts["ppt/_rels/presentation.xml.rels"].decode("utf-8")
    pres = parts["ppt/presentation.xml"].decode("utf-8")
    ids = []
    for i, (layout, shapes, charts) in enumerate(slides, start=1):
        rels = [f'<Relationship Id="rId1" Type="{REL_LAYOUT}" '
                f'Target="../slideLayouts/{Path(layout_file(base, layout)).name}"/>']
        for j, (chart, xml) in enumerate((charts or {}).items(), start=2):
            parts[f"ppt/charts/{chart}"] = xml.encode("utf-8")
            ct = ct.replace("</Types>", f'<Override PartName="/ppt/charts/{chart}" '
                                        f'ContentType="{CT_CHART}"/></Types>')
            rels.append(f'<Relationship Id="rId{j}" Type="{REL_CHART}" Target="../charts/{chart}"/>')
            shapes += (f'<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="{900 + j}" name="Chart {j}"/>'
                       f"<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr><p:xfrm/>"
                       f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
                       f'<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
                       f'r:id="rId{j}"/></a:graphicData></a:graphic></p:graphicFrame>')
        parts[f"ppt/slides/slide{i}.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<p:sld xmlns:a="{NS_A}" xmlns:r="{NS_R}" xmlns:p="{NS_P}"><p:cSld><p:spTree>'
            f'<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>'
            f"{shapes}</p:spTree></p:cSld></p:sld>").encode("utf-8")
        parts[f"ppt/slides/_rels/slide{i}.xml.rels"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            + "".join(rels) + "</Relationships>").encode("utf-8")
        ct = ct.replace("</Types>", f'<Override PartName="/ppt/slides/slide{i}.xml" '
                                    f'ContentType="{CT_SLIDE}"/></Types>')
        rid = f"rIdTest{i}"
        prels = prels.replace("</Relationships>", f'<Relationship Id="{rid}" Type="{REL_SLIDE}" '
                                                  f'Target="slides/slide{i}.xml"/></Relationships>')
        ids.append(f'<p:sldId id="{300 + i}" r:id="{rid}"/>')
    lst = "<p:sldIdLst>" + "".join(ids) + "</p:sldIdLst>"
    if "<p:sldIdLst/>" in pres:
        pres = pres.replace("<p:sldIdLst/>", lst)
    elif "<p:sldIdLst>" in pres:
        pres = pres.replace("<p:sldIdLst>", "<p:sldIdLst>" + "".join(ids), 1)
    else:
        pres = re.sub(r"(</p:sldMasterIdLst>)", r"\1" + lst, pres, count=1)
    parts["[Content_Types].xml"] = ct.encode("utf-8")
    parts["ppt/_rels/presentation.xml.rels"] = prels.encode("utf-8")
    parts["ppt/presentation.xml"] = pres.encode("utf-8")
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)
    return Path(dest)


def on_brand_slides():
    return [
        ("Standard TITLE", sp("title", None, "Remote digital towers") + sp("body", 2, "Ana Pop, 9 October 2026"), None),
        ("Sub-headline + field", sp("title", None, "Controllers see more with fewer screens")
         + sp("body", 14, "What changes for the tower team") + sp(None, 16, "One working position\nShared voice and data"), None),
        ("Closing Slide", "", None),
    ]


def minimal_docx(dest, body_runs, footer_text=None):
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    parts = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="xml" ContentType="application/xml"/><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        "word/document.xml": f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{w}"><w:body>'
        + "".join(f"<w:p><w:r><w:rPr>{rpr}</w:rPr><w:t>{t}</w:t></w:r></w:p>" for rpr, t in body_runs)
        + "</w:body></w:document>",
    }
    if footer_text:
        parts["word/footer1.xml"] = (f'<?xml version="1.0" encoding="UTF-8"?><w:ftr xmlns:w="{w}"><w:p><w:r>'
                                     f"<w:t>{footer_text}</w:t></w:r></w:p></w:ftr>")
    with zipfile.ZipFile(dest, "w") as z:
        for n, data in parts.items():
            z.writestr(n, data)
    return Path(dest)


def minimal_xlsx(dest, font="Calibri", rgb="FF000000", strings=("Budget",)):
    m = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    parts = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="xml" ContentType="application/xml"/></Types>',
        "xl/styles.xml": f'<?xml version="1.0" encoding="UTF-8"?><styleSheet xmlns="{m}"><fonts count="1"><font>'
        f'<sz val="11"/><color rgb="{rgb}"/><name val="{font}"/></font></fonts></styleSheet>',
        "xl/sharedStrings.xml": f'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="{m}">'
        + "".join(f"<si><t>{s}</t></si>" for s in strings) + "</sst>",
        "xl/worksheets/sheet1.xml": f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{m}"><sheetData>'
        '<row r="1"><c r="A1" t="s"><v>0</v></c></row></sheetData></worksheet>',
    }
    with zipfile.ZipFile(dest, "w") as z:
        for n, data in parts.items():
            z.writestr(n, data)
    return Path(dest)


def rules(found, severity=None, scope=None, location=None):
    return {f["rule"] for f in found
            if (severity is None or f["severity"] == severity)
            and (scope is None or f["scope"] == scope)
            and (location is None or f["location"] == location)}


class TestSkillFiles(unittest.TestCase):
    def test_frontmatter_and_triggers(self):
        text = (SKILL / "SKILL.md").read_text(encoding="utf-8")
        self.assertTrue(text.startswith("---\nname: frq-brandbook\n"))
        desc = re.search(r"(?m)^description: (.+)$", text).group(1)
        for word in ("Frequentis", "on-brand", "slides", "logo", "colours", "template", "check", "44", "spec",
                     "existing deck"):
            self.assertIn(word, desc)
        self.assertLessEqual(len(desc), 1024)
        self.assertLess(len(text.splitlines()), 220, "keep SKILL.md lean; move depth to references/")

    def test_quick_reference_matches_the_tokens(self):
        text = (SKILL / "SKILL.md").read_text(encoding="utf-8")
        tokens = json.loads(TOKENS.read_text(encoding="utf-8"))
        for c in tokens["colours"]["primary"] + tokens["colours"]["accent"]:
            self.assertIn(c["hex"], text, c["name"])
        for c in tokens["colours"]["primary"]:
            if "web_hex" in c:
                self.assertIn(c["web_hex"], text, c["name"])
        for cls in tokens["classification"]:
            self.assertIn(cls, text)

    def test_the_expected_files_are_there(self):
        for rel in ("SKILL.md", "PROVENANCE.md", "brand-tokens.json", "scripts/check_brand.py",
                    "scripts/new_deck.py", "assets/manifest.json",
                    "assets/templates/frq-template-slim-core.pptx",
                    "references/brand-rules.md", "references/building-decks.md",
                    "references/documents.md", "references/writing-style.md", "references/assets.md"):
            self.assertTrue((SKILL / rel).is_file(), rel)

    def test_the_skill_points_to_the_doc_skills_by_their_placed_names(self):
        text = (SKILL / "SKILL.md").read_text(encoding="utf-8")
        for skill in ("doc-powerpoint", "doc-word", "doc-excel", "doc-pdf", "visual-explainers",
                      "drawio", "likec4-dsl"):
            self.assertIn(f"`{packs.PREFIX}{skill}`", text)

    def test_every_file_of_the_skill_is_tracked_by_git(self):
        """An ignore rule (template/.gitignore has *.pptx) must not keep a skill file out of a
        clone; a dirty worktree would hide it, so ask git, not the disk."""
        r = subprocess.run(["git", "ls-files", "--", str(SKILL.relative_to(KIT))], cwd=KIT,
                           capture_output=True, text=True, check=False)
        if r.returncode != 0 or not (KIT / ".git").exists():
            self.skipTest("not a git checkout")
        tracked = {Path(x).relative_to(SKILL.relative_to(KIT)).as_posix() for x in r.stdout.split("\n") if x}
        on_disk = {p.relative_to(SKILL).as_posix() for p in SKILL.rglob("*")
                   if p.is_file() and "__pycache__" not in p.parts}
        self.assertEqual(sorted(on_disk - tracked), [])
        r = subprocess.run(["git", "check-ignore", "--no-index", "-q", str(TEMPLATE.relative_to(KIT))],
                           cwd=KIT, capture_output=True, check=False)
        self.assertEqual(r.returncode, 1, "the slim template is git-ignored")

    def test_logo_svgs_keep_their_proportions(self):
        for svg in sorted((SKILL / "assets/logo").glob("*.svg")):
            vb, wh = svg_box(svg)
            self.assertAlmostEqual(wh / vb, 1, delta=0.005, msg=svg.name)

    def test_logo_pngs_are_sharp_renders_of_the_svgs(self):
        for png in sorted((SKILL / "assets/logo").glob("*.png")):
            with self.subTest(png=png.name):
                w, h, mode, rows = png_alpha(png)
                self.assertEqual(mode, (8, 6, 0), "8-bit RGBA, not interlaced")
                vb, _ = svg_box(png.with_suffix(".svg"))
                self.assertAlmostEqual((w / h) / vb, 1, delta=0.006)     # same geometry as the SVG
                ink = [v for r in rows for v in r if v > 10]
                soft = sum(1 for v in ink if v < 245) / len(ink)
                self.assertLess(soft, 0.3, "soft edges: re-render the PNG from the SVG")

    def test_contrast_advice_holds(self):
        def lum(h):
            c = [int(h[i:i + 2], 16) / 255 for i in (1, 3, 5)]
            c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
            return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]

        def ratio(a, b):
            hi, lo = sorted((lum(a), lum(b)), reverse=True)
            return (hi + 0.05) / (lo + 0.05)
        self.assertLess(ratio("#FFFFFF", "#00AAE1"), 3)          # fails even for large text
        self.assertGreaterEqual(ratio("#333333", "#00AAE1"), 4.5)
        self.assertGreaterEqual(ratio("#004182", "#00AAE1"), 3)
        texts = [(SKILL / f).read_text(encoding="utf-8") for f in
                 ("SKILL.md", "brand-tokens.json", "references/brand-rules.md",
                  "references/documents.md", "references/building-decks.md")]
        for text in texts:
            for line in text.splitlines():
                if "white" in line.lower():
                    self.assertNotRegex(line, r"(?i)(18 ?pt|large text)[^\n]{0,30}(only|\+)", line)

    def test_brand_by_default_is_one_rule_in_each_file_skill(self):
        """Owner decision, 2026-10-09: with the brand skill installed, the file skills use the
        brand unless the person asks for a plain file."""
        for skill in ("doc-powerpoint", "doc-word", "doc-excel", "doc-pdf", "visual-explainers", "drawio",
                      "likec4-dsl"):
            text = (KIT / packs.SKILLS_REL / skill / "SKILL.md").read_text(encoding="utf-8")
            rules_ = [l for l in text.splitlines() if "ai-sdlc-frq-brandbook" in l]
            self.assertEqual(len(rules_), 1, skill)
            self.assertIn("Company brand by default", rules_[0], skill)
            self.assertRegex(rules_[0], r"unless the person asks for (a|the) plain", skill)
            # E2E 2026-10-09: Copilot's search skips the git-excluded placed folder, so it
            # concluded the brand skill was missing. Read by exact path, never by search.
            self.assertIn("If it is in your skill list, it is installed", rules_[0], skill)
            self.assertIn("`.agents/skills/ai-sdlc-frq-brandbook/…`", rules_[0], skill)
            self.assertIn("Never decide by glob or search", rules_[0], skill)
        desc = re.search(r"(?m)^description: (.+)$", (SKILL / "SKILL.md").read_text(encoding="utf-8")).group(1)
        self.assertIn("default", desc)
        for word in ("branded", "our template", "our colours"):
            self.assertIn(word, desc)

    def test_the_classification_is_asked_and_never_guessed_in_every_recipe(self):
        """E2E 2026-10-09: decks, explainers and footers got a guessed class."""
        skill = (SKILL / "SKILL.md").read_text(encoding="utf-8")
        for needed in ("Never guess it", "`Frequentis [classification to be set]`",
                       "Outline first, then stop", "Build nothing until the person agrees",
                       "Fix every FAIL and every font WARN", "never import them"):
            self.assertIn(needed, skill)
        docs = (SKILL / "references/documents.md").read_text(encoding="utf-8")
        self.assertIn("**Classification: ask, never guess.**", docs)
        self.assertIn("`Frequentis [classification to be set]`", docs)
        self.assertEqual(json.loads(TOKENS.read_text(encoding="utf-8"))["classification_placeholder"],
                         "Frequentis [classification to be set]")
        decks = (SKILL / "references/building-decks.md").read_text(encoding="utf-8")
        self.assertNotRegex(decks + skill, r'--classification "Frequentis (Public|General|Confidential)"',
                            "an example with a real class gets copied as is")

    def test_render_step_is_one_fixed_command_into_the_hidden_tmp_folder(self):
        skill = (SKILL / "SKILL.md").read_text(encoding="utf-8")
        self.assertIn("`soffice --headless --convert-to pdf --outdir .ai-sdlc/tmp <file>`", skill)
        self.assertIn("`pdftoppm -png -r 80 .ai-sdlc/tmp/<name>.pdf .ai-sdlc/tmp/<name>`", skill)

    def test_the_font_snippets_come_first_in_the_word_and_excel_recipes(self):
        docs = (SKILL / "references/documents.md").read_text(encoding="utf-8")
        for head, snippet in (("## Word (with", "w:asciiTheme"), ("## Excel (with", "def arial_everywhere")):
            section = docs.split(head, 1)[1]
            first = section.split("\n\n", 2)[1]
            self.assertIn("first" if "Word" in head else "every openpyxl script", first, head)
            self.assertLess(section.index(snippet), section.index("\n- "), head)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is not installed here")
    def test_the_excel_recipe_leaves_no_font_warning(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        docs = (SKILL / "references/documents.md").read_text(encoding="utf-8")
        recipe = docs.split("## Excel (with", 1)[1].split("```python", 1)[1].split("```", 1)[0]
        ns = {}
        exec(compile(recipe, "documents.md", "exec"), ns)
        with tempfile.TemporaryDirectory() as tmp:
            wb = Workbook()
            ws = wb.active
            ws.append(["Team", "Days"])
            ws.append(["Ana", 9])
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="004182")
            ws.oddFooter.left.text = "Frequentis General | © Frequentis AG 2026"
            ns["arial_everywhere"](wb)
            out = Path(tmp) / "x.xlsx"
            wb.save(str(out))
            code, found = findings(out, "--year", "2026")
            self.assertEqual([f for f in found if f["rule"] == "font.non-brand"], [])
            self.assertTrue(Font is not None and ws["A1"].font.b)

    def test_customer_word_work_points_to_the_official_template(self):
        for rel in ("SKILL.md", "references/documents.md"):
            self.assertIn("Doknorme.dotm", (SKILL / rel).read_text(encoding="utf-8"), rel)
        self.assertNotIn("no official Word", (SKILL / "references/documents.md").read_text(encoding="utf-8"))

    def test_size_budget(self):
        total = sum(p.stat().st_size for p in SKILL.rglob("*") if p.is_file())
        self.assertLessEqual(total, BUDGET_ALL, f"{total} bytes")

    def test_manifest_lists_every_asset_with_its_size(self):
        man = json.loads((SKILL / "assets/manifest.json").read_text(encoding="utf-8"))
        on_disk = {p.relative_to(SKILL).as_posix(): p.stat().st_size
                   for p in (SKILL / "assets").rglob("*") if p.is_file() and p.name != "manifest.json"}
        listed = {a["file"]: a for a in man["assets"]}
        self.assertEqual(sorted(listed), sorted(on_disk))
        for rel, a in listed.items():
            self.assertEqual(a["bytes"], on_disk[rel], rel)
            self.assertEqual(a["sha256"], hashlib.sha256((SKILL / rel).read_bytes()).hexdigest(), rel)
            for key in ("file", "type", "variant", "bytes", "use"):
                self.assertIn(key, a, rel)

    def test_nothing_with_unknown_rights_is_bundled(self):
        for p in SKILL.rglob("*"):
            rel = p.relative_to(SKILL).as_posix().lower()
            self.assertNotRegex(rel, r"\.(ttf|otf|woff2?|eot)$", "never bundle font files")
            self.assertNotRegex(rel, r"(^|/)photo|portrait|sesar|fabec|screenshot|stock", rel)

    def test_the_template_is_slim_and_carries_no_personal_metadata(self):
        with zipfile.ZipFile(TEMPLATE) as z:
            names = z.namelist()
            self.assertFalse([n for n in names if n.startswith("ppt/slides/")], "no sample slides")
            layouts = [n for n in names if re.fullmatch(r"ppt/slideLayouts/slideLayout\d+\.xml", n)]
            self.assertEqual(len(layouts), 25)
            for gone in ("ppt/commentAuthors.xml", "docProps/custom.xml", "docProps/thumbnail.jpeg"):
                self.assertNotIn(gone, names)
            self.assertFalse([n for n in names if n.startswith("customXml/")])
            core = z.read("docProps/core.xml").decode("utf-8")
            self.assertIn("<dc:creator></dc:creator>", core)
            for n in names:
                if n.endswith((".xml", ".rels")):
                    self.assertNotIn("@", z.read(n).decode("utf-8").replace("@ ", ""), n)

    def test_provenance_and_name_screen(self):
        prov = (SKILL / "PROVENANCE.md").read_text(encoding="utf-8")
        for topic in ("Brand Guidelines Q4/2025", "shape for shape",
                      "Group Communications", "Icon Stock", "C1", "C2", "C6"):
            self.assertIn(topic, prov)
        # Built from pieces, so this file does not match its own screen.
        words = ["/" + "Users/", "~" + "/work", "emi" + "dan", "Publi" + "cis", "Sap" + "ient",
                 "ps-" + "brandbook", "NOTES" + "/"]
        banned = re.compile("|".join(re.escape(w) for w in words), re.I)
        files = [p for p in SKILL.rglob("*") if p.is_file()] + [Path(__file__)]
        for p in files:
            try:
                text = p.read_bytes().decode("utf-8")
            except UnicodeDecodeError:
                continue
            with self.subTest(path=p.name):
                self.assertIsNone(banned.search(text))

    def test_tokens_are_consistent(self):
        tokens = json.loads(TOKENS.read_text(encoding="utf-8"))
        hexes = [c["hex"] for group in ("primary", "accent", "legend_only")
                 for c in tokens["colours"][group]]
        self.assertEqual(len(hexes), len(set(hexes)))
        for h in hexes:
            self.assertRegex(h, r"^#[0-9A-F]{6}$")
        for c in tokens["colours"]["primary"] + tokens["colours"]["accent"]:
            self.assertEqual("#%02X%02X%02X" % tuple(c["rgb"]), c["hex"], c["name"])
        self.assertEqual(tokens["fonts"]["office"]["family"], "Arial")
        for bu in tokens["business_units"]:
            for key in ("key_visual", "key_visual_full", "key_visual_wide"):
                if key in bu:
                    self.assertTrue((SKILL / bu[key]).is_file(), (bu["id"], key))


def inventory():
    return json.loads(INVENTORY.read_text(encoding="utf-8"))


def kit_only_patterns():
    return [line.strip() for line in KIT_ONLY.read_text(encoding="utf-8").splitlines()
            if line.strip() and not line.lstrip().startswith("#")]


def is_kit_only(rel):
    """Matched by .kit-only, as setup reads it (place.kit_only)."""
    return rel in set(place.kit_only(KIT, "frq-brandbook"))


def skill_files():
    return sorted(p.relative_to(SKILL).as_posix() for p in SKILL.rglob("*")
                  if p.is_file() and "__pycache__" not in p.parts)


def is_text(rel):
    try:
        (SKILL / rel).read_bytes().decode("utf-8")
        return True
    except UnicodeDecodeError:
        return False


# Text files of the owner skill and where their content lives now (design §8.2). The master is
# here too: it is a binary, but cleaned (B1), so it is mapped rather than byte-identical.
MAPPED_TEXT = {"SKILL.md": "SKILL.md", "references/brand-rules.md": "references/brand-rules.md",
               "references/layouts.md": "references/layouts.md",
               "references/build-spec.md": "references/build-spec.md",
               "scripts/frq_pptx.py": "scripts/frq_pptx.py", "assets/brand-tokens.json": "brand-tokens.json",
               "assets/frequentis-brand.css": "assets/frequentis-brand.css",
               "assets/frq-master.pptx": "assets/templates/frq-master.pptx"}
DUPLICATES = {"assets/key-visuals/atm-aircraft.jpg": "assets/keyvisual/keyvisual-atm-aircraft.jpeg",
              "assets/key-visuals/defence.jpg": "assets/keyvisual/keyvisual-defence-jets.jpeg",
              "assets/key-visuals/maritime.jpg": "assets/keyvisual/keyvisual-maritime-vessel.jpeg",
              "assets/key-visuals/public-safety.jpg": "assets/keyvisual/keyvisual-public-safety-police.jpeg"}


class TestMergedAssets(unittest.TestCase):
    """The owner skill's previews, examples, key visuals, logos, CSS and tokens: nothing lost
    (design §8.2), and every binary kept only in the kit copy (design §8.7)."""

    def setUp(self):
        self.manifest = {a["file"]: a for a in
                         json.loads((SKILL / "assets/manifest.json").read_text(encoding="utf-8"))["assets"]}

    def test_every_source_binary_is_in_the_skill_byte_for_byte(self):
        here = {hashlib.sha256((SKILL / rel).read_bytes()).hexdigest() for rel in skill_files()}
        for f in inventory()["files"]:
            if f["path"] in MAPPED_TEXT:
                continue
            with self.subTest(path=f["path"]):
                self.assertIn(f["sha256"], here)

    def test_every_layout_preview_and_example_is_present_and_listed(self):
        for folder, count in (("assets/layouts", 44), ("assets/examples", 22)):
            files = sorted(p.relative_to(SKILL).as_posix() for p in (SKILL / folder).glob("*.jpg"))
            self.assertEqual(len(files), count, folder)
            for rel in files:
                a = self.manifest[rel]
                self.assertEqual(a["bytes"], (SKILL / rel).stat().st_size, rel)
                self.assertEqual(a["source"], "owner skill v1.0", rel)
                self.assertEqual(a["from"], rel, "names unchanged")

    def test_duplicates_are_kept_once_with_an_alias(self):
        inv = {f["path"]: f["sha256"] for f in inventory()["files"]}
        for owner, ours in DUPLICATES.items():
            self.assertEqual(hashlib.sha256((SKILL / ours).read_bytes()).hexdigest(), inv[owner], owner)
            self.assertEqual(self.manifest[ours]["aliases"], [owner])
            self.assertEqual(self.manifest[ours]["source"], "0.8.0")
            self.assertFalse((SKILL / "assets/keyvisual" / Path(owner).name).exists(), "kept once")
        for owner, ours in (("assets/key-visuals/corporate-globe.jpg", "keyvisual-corporate-globe.jpg"),
                            ("assets/key-visuals/public-transport.jpg", "keyvisual-public-transport.jpg"),
                            ("assets/key-visuals/atm-aircraft-clouds-wide.jpg",
                             "keyvisual-atm-aircraft-clouds-wide.jpg")):
            self.assertEqual(self.manifest["assets/keyvisual/" + ours]["from"], owner)

    def test_both_logo_sets_and_the_converted_ones_are_preferred(self):
        for colour in ("blue", "white", "black"):
            self.assertIs(self.manifest[f"assets/logo/frequentis-logo-{colour}.svg"]["preferred"], False)
            self.assertIs(self.manifest[f"assets/logo/logo-frequentis-wordmark-{colour}.svg"]["preferred"], True)

    def test_tokens_and_css_agree(self):
        tokens = json.loads(TOKENS.read_text(encoding="utf-8"))
        col = tokens["colours"]
        hexes = {c[k].upper() for g in ("primary", "accent", "legend_only", "chart_only") for c in col[g]
                 for k in ("hex", "web_hex") if k in c}
        css = (SKILL / "assets/frequentis-brand.css").read_text(encoding="utf-8")
        found = {h.upper() for h in re.findall(r"#[0-9A-Fa-f]{6}\b", css)}
        self.assertTrue(found)
        self.assertLessEqual(found, hexes)
        self.assertIn(col["gradient"]["from"], css)
        self.assertIn(col["gradient"]["to"], css)
        self.assertEqual(tokens["charts"]["series"], col["chart_series_order"], "one series order (C17)")
        self.assertEqual(tokens["charts"]["series"][:2], ["#004182", "#00AAE1"])
        self.assertEqual(tokens["charts"]["kpi_track"], "#EDF1F2")
        self.assertEqual([c["hex"] for c in col["chart_only"]], ["#EDF1F2"])
        self.assertEqual(tokens["pptx"]["template_full"], "assets/templates/frq-master.pptx")
        self.assertEqual(tokens["pptx"]["content_area_in"]["bottom"], 5.07)
        self.assertIn("<Presenter>", tokens["footer_format_full"])
        self.assertEqual(tokens["default_business_unit"], "atm")

    def test_size_budget_placed(self):
        placed = sum((SKILL / rel).stat().st_size for rel in skill_files()
                     if not is_kit_only(rel) and not rel.startswith("."))
        self.assertLessEqual(placed, BUDGET_PLACED, f"{placed} bytes placed")

    def test_every_binary_file_is_kit_only(self):
        files = skill_files()
        for rel in files:
            if not is_text(rel):
                self.assertTrue(is_kit_only(rel), rel)
        self.assertTrue(kit_only_patterns())
        self.assertEqual([e for e in place.validate_kit(KIT) if ".kit-only" in e], [], "every pattern matches")
        self.assertIn("design 2026-10-09 §8.7", KIT_ONLY.read_text(encoding="utf-8"))

    def test_the_manifest_says_where_each_asset_lives(self):
        for rel, a in self.manifest.items():
            self.assertEqual(a["placement"], "kit-only" if is_kit_only(rel) else "placed", rel)
            self.assertIn(a["source"], ("0.8.0", "owner skill v1.0"), rel)


class TestFullMaster(unittest.TestCase):
    """The full official master, cleaned of personal and tenant data (design §8.3)."""

    def test_the_inventory_lists_the_84_files_of_the_owner_skill(self):
        inv = inventory()
        self.assertEqual((inv["skill"], inv["version"], inv["date"]), ("frq-4-pptx-agent", "1.0", "2026-10-09"))
        paths_ = [f["path"] for f in inv["files"]]
        self.assertEqual(len(paths_), 84)
        self.assertEqual(paths_, sorted(paths_))
        for f in inv["files"]:
            self.assertFalse(f["path"].startswith("/") or ".." in f["path"], f["path"])
            self.assertRegex(f["sha256"], r"^[0-9a-f]{64}$")
        self.assertNotIn("@", INVENTORY.read_text(encoding="utf-8"))

    def test_the_full_master_has_44_layouts_and_no_slides(self):
        with zipfile.ZipFile(FULL) as z:
            names = z.namelist()
        self.assertFalse([n for n in names if n.startswith("ppt/slides/")], "no sample slides")
        self.assertEqual(len([n for n in names if re.fullmatch(r"ppt/slideLayouts/slideLayout\d+\.xml", n)]), 44)
        self.assertEqual(len([n for n in names if re.fullmatch(r"ppt/slideMasters/slideMaster\d+\.xml", n)]), 1)

    def test_the_full_master_carries_no_personal_or_tenant_metadata(self):
        with zipfile.ZipFile(FULL) as z:
            names = z.namelist()
            for gone in ("ppt/commentAuthors.xml", "docProps/custom.xml", "docProps/thumbnail.jpeg"):
                self.assertNotIn(gone, names)
            self.assertFalse([n for n in names if n.startswith("customXml/")])
            core = z.read("docProps/core.xml").decode("utf-8")
            self.assertIn("<dc:creator></dc:creator>", core)
            self.assertIn("<cp:lastModifiedBy></cp:lastModifiedBy>", core)
            app = z.read("docProps/app.xml").decode("utf-8")
            self.assertEqual(re.findall(r"<(\w+)>", app), ["Application", "PresentationFormat"])
            for n in names:
                if n.endswith((".xml", ".rels")):
                    text = z.read(n).decode("utf-8")
                    self.assertNotIn("@", text.replace("@ ", ""), n)
                    self.assertNotIn("MSIP_", text, n)
                    for gone in ("commentAuthors", "customXml", "thumbnail", "custom-properties"):
                        self.assertNotIn(gone, text, n)

    def test_only_metadata_changed(self):
        parts = inventory()["master_parts"]
        self.assertTrue([n for n in parts if n.startswith("ppt/slideLayouts/")])
        self.assertTrue([n for n in parts if n.startswith("ppt/media/")])
        with zipfile.ZipFile(FULL) as z:
            here = {n for n in z.namelist() if n.startswith(("ppt/slideMasters/", "ppt/slideLayouts/",
                                                              "ppt/theme/", "ppt/media/"))}
            self.assertEqual(sorted(here), sorted(parts))
            for n, sha in parts.items():
                self.assertEqual(hashlib.sha256(z.read(n)).hexdigest(), sha, n)


class TestCheckBrandScript(unittest.TestCase):
    def test_stdlib_only_and_python_3_9(self):
        src = CHECK.read_text(encoding="utf-8")
        tree = ast.parse(src, feature_version=(3, 9))
        allowed = {"__future__", "argparse", "datetime", "importlib", "json", "re", "sys", "zipfile", "xml",
                   "pathlib", "posixpath", "os", "collections", "zlib"}
        for script in (CHECK, ASSETS):          # check_brand.py loads brand_assets.py by path
            tree = ast.parse(script.read_text(encoding="utf-8"), feature_version=(3, 9))
            mods = set()
            for node in ast.walk(tree):
                if isinstance(node, ast.Import):
                    mods |= {a.name.split(".")[0] for a in node.names}
                elif isinstance(node, ast.ImportFrom):
                    mods.add((node.module or "").split(".")[0])
            self.assertLessEqual(mods, allowed, script.name)

    def test_golden_off_brand_deck(self):
        code, found = findings(GOLDEN, "--year", "2026")
        self.assertEqual(code, 1)
        fail = {loc: rules(found, "FAIL", "slide", loc) for loc in
                ("slide 1", "slide 2", "slide 3", "slide 4", "slide 5", "slide 6")}
        self.assertLessEqual({"font.non-brand", "text.italic", "colour.off-palette", "text.frq"},
                             fail["slide 1"])
        self.assertLessEqual({"text.tagline", "text.italic", "colour.off-palette"}, fail["slide 2"])
        self.assertLessEqual({"font.non-brand", "colour.off-palette", "effect.shadow",
                              "gradient.non-brand"}, fail["slide 3"])
        self.assertLessEqual({"colour.off-palette", "text.italic"}, fail["slide 4"])
        self.assertLessEqual({"gradient.non-brand", "colour.off-palette"}, fail["slide 5"])
        self.assertIn("slide.thank-you", fail["slide 6"])
        warn = {loc: rules(found, "WARN", "slide", loc) for loc in fail}
        self.assertLessEqual({"text.align", "text.exclamation", "text.caps-name", "text.date"},
                             warn["slide 1"])
        self.assertLessEqual({"text.title-case", "text.us-spelling", "text.bold", "text.align",
                              "text.at-frequentis", "text.contraction", "text.number"},
                             warn["slide 2"])
        self.assertIn("shape.outline", warn["slide 3"])
        self.assertIn("text.us-spelling", warn["slide 4"])              # chart series "Defense"
        self.assertLessEqual({"slide.bullets", "text.title-case"}, warn["slide 5"])
        self.assertIn("footer.classification", rules(found, "FAIL", "file"))
        self.assertLessEqual({"theme.fonts", "theme.colours"}, rules(found, "WARN", "theme"))
        self.assertIn("text.ampersand", rules(found, "INFO", "slide", "slide 1"))

    def test_findings_carry_a_fix_and_the_text_report_is_readable(self):
        code, found = findings(GOLDEN)
        for f in found:
            self.assertTrue(f["message"] and f["fix"] and f["rule"], f)
            self.assertIn(f["severity"], ("FAIL", "WARN", "INFO"))
        code, out, _ = run_check(GOLDEN)
        self.assertEqual(code, 1)
        self.assertIn("FAIL", out)
        self.assertRegex(out, r"\d+ FAIL, \d+ WARN, \d+ INFO")

    def test_findings_are_numbered_so_a_person_can_pick_by_number(self):
        code, found = findings(GOLDEN)
        self.assertEqual([f["id"] for f in found], list(range(1, len(found) + 1)))
        _, out, _ = run_check(GOLDEN)
        self.assertIn("| # | Severity | Where |", out)
        rows = [l for l in out.splitlines() if re.match(r"^\| \d+ \| (FAIL|WARN|INFO) \|", l)]
        self.assertEqual([int(r.split("|")[1]) for r in rows], list(range(1, len(found) + 1)))

    def test_the_template_note_is_there_only_with_template_findings(self):
        note = "come from the template"
        with tempfile.TemporaryDirectory() as tmp:
            deck = build_pptx(Path(tmp) / "ok.pptx", on_brand_slides())
            _, out, _ = run_check(deck)
            self.assertIn(note, out)
            xlsx = minimal_xlsx(Path(tmp) / "x.xlsx", font="Arial", rgb="FF004182",
                                strings=("Budget", "Frequentis General"))
            _, out, _ = run_check(xlsx)
            self.assertNotIn(note, out)

    def test_a_classification_still_to_be_set_is_a_warning_not_a_guess(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = minimal_docx(Path(tmp) / "d.docx", [
                ('<w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>', "Report")],
                footer_text="Frequentis [classification to be set] | © Frequentis AG 2026")
            code, found = findings(doc, "--year", "2026")
            self.assertEqual(code, 0)
            self.assertIn("footer.classification-to-set", rules(found, "WARN", "file"))
            self.assertNotIn("footer.classification", rules(found))

    def test_the_scripts_leave_no_bytecode_behind(self):
        """E2E 2026-10-09: a __pycache__ in the placed skill showed up in `check`."""
        for script in (CHECK, NEW_DECK, FRQ_PPTX, ASSETS):
            src = script.read_text(encoding="utf-8")
            self.assertIn("\nsys.dont_write_bytecode = True", src, script.name)
            self.assertLess(src.index("sys.dont_write_bytecode = True"),
                            src.index("from pathlib import" if script == ASSETS else "import argparse"))

    def test_a_deck_built_from_the_template_passes(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = build_pptx(Path(tmp) / "ok.pptx", on_brand_slides())
            code, found = findings(deck, "--year", "2024")     # the template master says 2024
            self.assertEqual([f for f in found if f["severity"] != "INFO"], [])
            self.assertEqual(code, 0)
            # master and layout styles are reported, but only as INFO (conflicts C2, C3, C9)
            self.assertTrue([f for f in found if f["scope"] in ("master", "layout")])
            self.assertIn("footer.classification", rules(found, "INFO", "file"))
            code, found = findings(deck, "--year", "2026")
            self.assertEqual(code, 0)
            self.assertEqual({(f["severity"], f["rule"]) for f in found if f["severity"] != "INFO"},
                             {("WARN", "footer.year")})

    def test_slide_level_overrides_on_the_template_fail(self):
        bad = ('<a:rPr i="1"><a:solidFill><a:srgbClr val="FF00FF"/></a:solidFill>'
               '<a:latin typeface="Calibri"/></a:rPr>')
        shadow = ('<a:solidFill><a:srgbClr val="004182"/></a:solidFill>'
                  '<a:effectLst><a:outerShdw blurRad="1"><a:srgbClr val="000000"/></a:outerShdw></a:effectLst>')
        chart = ('<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
                 f'xmlns:a="{NS_A}"><c:chart><c:view3D/><c:plotArea><c:bar3DChart/></c:plotArea>'
                 "</c:chart></c:chartSpace>")
        with tempfile.TemporaryDirectory() as tmp:
            deck = build_pptx(Path(tmp) / "bad.pptx", [
                ("Headline (standard)", sp("title", None, "FRQ results", rpr=bad), None),
                ("Headline (standard)", sp(None, None, "Box", sppr=shadow), {"chart1.xml": chart}),
                ("Headline (standard)", sp("title", None, "Thank you"), None)])
            code, found = findings(deck, "--year", "2024")
            self.assertEqual(code, 1)
            self.assertLessEqual({"text.italic", "colour.off-palette", "font.non-brand"},
                                 rules(found, "FAIL", "slide", "slide 1"))
            self.assertIn("text.frq", rules(found, "WARN", "slide", "slide 1"))   # class: General
            self.assertLessEqual({"effect.shadow", "effect.3d"}, rules(found, "FAIL", "slide", "slide 2"))
            self.assertNotIn("colour.off-palette", rules(found, "FAIL", "slide", "slide 2"),
                             "a shadow's colour is the shadow finding, not a second one")
            self.assertIn("slide.thank-you", rules(found, "FAIL", "slide", "slide 3"))

    def test_theme_styles_that_python_pptx_shapes_inherit_are_seen(self):
        style = ('<p:style><a:lnRef idx="1"><a:schemeClr val="accent1"><a:shade val="50000"/></a:schemeClr></a:lnRef>'
                 '<a:fillRef idx="3"><a:schemeClr val="accent1"/></a:fillRef>'
                 '<a:effectRef idx="2"><a:schemeClr val="accent1"/></a:effectRef>'
                 '<a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef></p:style>')
        default = ('<p:sp><p:nvSpPr><p:cNvPr id="40" name="Rectangle 1"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
                   '<p:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>' + style + '</p:sp>')
        clean = ('<p:sp><p:nvSpPr><p:cNvPr id="41" name="Rectangle 2"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
                 '<p:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="004182"/>'
                 '</a:solidFill><a:ln><a:noFill/></a:ln><a:effectLst/></p:spPr>' + style + '</p:sp>')
        with tempfile.TemporaryDirectory() as tmp:
            deck = build_pptx(Path(tmp) / "shapes.pptx", [("Headline (standard)", default, None),
                                                         ("Headline (standard)", clean, None)])
            code, found = findings(deck, "--year", "2024")
            self.assertEqual(code, 1)
            self.assertLessEqual({"effect.shadow", "gradient.non-brand"}, rules(found, "FAIL", "slide", "slide 1"))
            self.assertIn("shape.outline", rules(found, "WARN", "slide", "slide 1"))
            self.assertEqual(rules(found, None, "slide", "slide 2"), set())

    def test_tints_and_system_colours(self):
        tint = '<a:rPr><a:solidFill><a:schemeClr val="accent3"><a:lumMod val="60000"/></a:schemeClr></a:solidFill></a:rPr>'
        grey = '<a:rPr><a:solidFill><a:schemeClr val="bg2"><a:lumMod val="60000"/></a:schemeClr></a:solidFill></a:rPr>'
        sysc = '<a:rPr><a:solidFill><a:sysClr val="windowText" lastClr="000000"/></a:solidFill></a:rPr>'
        prst = '<a:rPr><a:solidFill><a:prstClr val="black"/></a:solidFill></a:rPr>'
        with tempfile.TemporaryDirectory() as tmp:
            deck = build_pptx(Path(tmp) / "tints.pptx", [
                ("Headline (standard)", sp(None, None, "Light blue tint", rpr=tint), None),
                ("Headline (standard)", sp(None, None, "Grey tint", rpr=grey), None),
                ("Headline (standard)", sp(None, None, "System black", rpr=sysc), None),
                ("Headline (standard)", sp(None, None, "Preset black", rpr=prst), None)])
            code, found = findings(deck, "--year", "2024")
            self.assertIn("colour.tint", rules(found, "WARN", "slide", "slide 1"))
            self.assertEqual(rules(found, None, "slide", "slide 2"), set())
            self.assertIn("colour.off-palette", rules(found, "FAIL", "slide", "slide 3"))
            self.assertIn("colour.off-palette", rules(found, "FAIL", "slide", "slide 4"))

    def test_name_and_tagline_severity_follow_the_rules(self):
        with tempfile.TemporaryDirectory() as tmp:
            internal = build_pptx(Path(tmp) / "internal.pptx", [   # the master says Frequentis General
                ("Headline (standard)", sp("title", None, "FRQ status for the team")
                 + sp(None, None, "FREQUENTIS COMSOFT GmbH supplies it.\nWe work for a safer world every day."), None),
                ("Headline (standard)", sp(None, None, "FREQUENTIS", name="Logo text"), None),
                ("Headline (standard)", sp(None, None, "For a safer world", name="Tagline"), None)])
            code, found = findings(internal, "--year", "2024")
            self.assertIn("text.frq", rules(found, "WARN", "slide", "slide 1"))
            self.assertNotIn("text.frq", rules(found, "FAIL"))
            self.assertNotIn("text.caps-name", rules(found, None, "slide", "slide 1"))
            self.assertIn("text.tagline", rules(found, "WARN", "slide", "slide 1"))
            self.assertIn("text.caps-name", rules(found, "WARN", "slide", "slide 2"))
            self.assertIn("text.tagline", rules(found, "FAIL", "slide", "slide 3"))
            public = build_pptx(Path(tmp) / "public.pptx", [
                ("Headline (standard)", sp("title", None, "FRQ results") + sp(None, None, "Frequentis Public"), None)])
            self.assertIn("text.frq", rules(findings(public, "--year", "2024")[1], "FAIL", "slide", "slide 1"))

    def test_a_questions_closing_slide_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = build_pptx(Path(tmp) / "q.pptx", [("Headline (standard)", sp("title", None, "Questions?"), None)])
            self.assertIn("slide.thank-you", rules(findings(deck, "--year", "2024")[1], "FAIL", "slide", "slide 1"))

    def test_word_checks_used_styles_and_theme_fonts_only(self):
        w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        styles = (f'<?xml version="1.0" encoding="UTF-8"?><w:styles xmlns:w="{w}">'
                  '<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:asciiTheme="minorHAnsi" w:hAnsiTheme="minorHAnsi"/>'
                  '</w:rPr></w:rPrDefault></w:docDefaults>'
                  '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial"/></w:rPr></w:style>'
                  '<w:style w:type="paragraph" w:styleId="Heading1"><w:basedOn w:val="Normal"/><w:rPr>'
                  '<w:rFonts w:asciiTheme="majorHAnsi" w:hAnsiTheme="majorHAnsi"/><w:color w:val="004182"/></w:rPr></w:style>'
                  '<w:style w:type="character" w:styleId="Code"><w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New"/>'
                  '<w:i/><w:color w:val="FF00FF"/></w:rPr></w:style></w:styles>')
        theme = ('<?xml version="1.0" encoding="UTF-8"?><a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                 '<a:themeElements><a:clrScheme name="Office"><a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>'
                 '<a:accent1><a:srgbClr val="4472C4"/></a:accent1></a:clrScheme><a:fontScheme name="Office">'
                 '<a:majorFont><a:latin typeface="Calibri Light"/></a:majorFont><a:minorFont><a:latin typeface="Arial"/></a:minorFont>'
                 '</a:fontScheme></a:themeElements></a:theme>')
        with tempfile.TemporaryDirectory() as tmp:
            path = minimal_docx(Path(tmp) / "doc.docx", [("", "Frequentis report")],
                                footer_text="Frequentis General | © Frequentis AG 2026")
            with zipfile.ZipFile(path, "a") as z:
                z.writestr("word/styles.xml", styles)
                z.writestr("word/theme/theme1.xml", theme)
            body = zipfile.ZipFile(path).read("word/document.xml").decode().replace(
                "<w:p><w:r>", '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r>', 1)
            parts = {n: zipfile.ZipFile(path).read(n) for n in zipfile.ZipFile(path).namelist()}
            parts["word/document.xml"] = body.encode()
            with zipfile.ZipFile(path, "w") as z:
                for n, d in parts.items():
                    z.writestr(n, d)
            code, found = findings(path, "--year", "2026")
            fonts = [f for f in found if f["rule"] == "font.non-brand"]
            self.assertTrue(fonts and "Calibri Light" in fonts[0]["message"], found)
            self.assertNotIn("Courier New", json.dumps(found))           # unused style: not reported
            self.assertNotIn("text.italic", rules(found))
            self.assertLessEqual(rules(found, None, "theme"), rules(found, "INFO", "theme"))

    def test_word_and_excel(self):
        with tempfile.TemporaryDirectory() as tmp:
            bad = minimal_docx(Path(tmp) / "bad.docx", [
                ('<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:i/><w:color w:val="FF0000"/>', "FRQ report")])
            code, found = findings(bad)
            self.assertEqual(code, 1)
            self.assertLessEqual({"font.non-brand", "text.italic", "colour.off-palette", "text.frq"},
                                 rules(found, "FAIL"))
            self.assertIn("footer.classification", rules(found, "FAIL", "file"))
            good = minimal_docx(Path(tmp) / "good.docx", [
                ('<w:rFonts w:ascii="Arial" w:hAnsi="Arial"/><w:color w:val="004182"/>', "Frequentis report")],
                footer_text="Frequentis General | © Frequentis AG 2026")
            code, found = findings(good, "--year", "2026")
            self.assertEqual((code, rules(found, "FAIL")), (0, set()))
            xbad = minimal_xlsx(Path(tmp) / "bad.xlsx", strings=("FRQ budget",))
            code, found = findings(xbad)
            self.assertEqual(code, 1)
            self.assertIn("text.frq", rules(found, "FAIL"))
            self.assertIn("font.non-brand", rules(found, "WARN"))      # Excel default styles: WARN
            xgood = minimal_xlsx(Path(tmp) / "good.xlsx", font="Arial", rgb="FF004182",
                                 strings=("Budget", "Frequentis General"))
            self.assertEqual(findings(xgood)[0], 0)

    def test_errors_and_usage(self):
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "x.txt").write_text("hi")
            self.assertEqual(run_check(Path(tmp) / "x.txt")[0], 2)
            (Path(tmp) / "broken.pptx").write_bytes(b"not a zip")
            code, _, err = run_check(Path(tmp) / "broken.pptx")
            self.assertEqual(code, 2)
            self.assertIn("broken.pptx", err)
            self.assertEqual(run_check(Path(tmp) / "missing.pptx")[0], 2)

    def malformed(self, tmp, name, slide_bytes, compress=zipfile.ZIP_DEFLATED):
        """A copy of the golden deck whose slide1.xml is replaced by `slide_bytes`."""
        dest = Path(tmp) / name
        with zipfile.ZipFile(GOLDEN) as zin, zipfile.ZipFile(dest, "w", compress) as zout:
            for n in zin.namelist():
                zout.writestr(n, slide_bytes if n == "ppt/slides/slide1.xml" else zin.read(n))
        return dest

    def assert_refused(self, path, words):
        code, out, err = run_check(path)
        self.assertEqual(code, 2, out + err)
        self.assertNotIn("Traceback", err)
        self.assertRegex(err, words)

    def test_entity_expansion_is_refused_in_any_encoding(self):
        lol = ('<?xml version="1.0"?><!DOCTYPE l [<!ENTITY a "FRQ FRQ FRQ FRQ">'
               '<!ENTITY b "&a;&a;&a;&a;&a;&a;"><!ENTITY c "&b;&b;&b;&b;&b;&b;">]>'
               f'<p:sld xmlns:p="{NS_P}" xmlns:a="{NS_A}"><a:t>&c;</a:t></p:sld>')
        with tempfile.TemporaryDirectory() as tmp:
            self.assert_refused(self.malformed(tmp, "lol8.pptx", lol.encode("utf-8")), "DTD|entit")
            utf16 = lol.replace('version="1.0"', 'version="1.0" encoding="UTF-16"').encode("utf-16")
            self.assert_refused(self.malformed(tmp, "lol16.pptx", utf16), "UTF-8")
            nobom = lol.encode("utf-16-le")                        # no BOM: NUL bytes give it away
            self.assert_refused(self.malformed(tmp, "lol16le.pptx", nobom), "UTF-8")
            latin = ('<?xml version="1.0" encoding="ISO-8859-1"?>'
                     f'<p:sld xmlns:p="{NS_P}"/>').encode("latin-1")
            self.assert_refused(self.malformed(tmp, "latin.pptx", latin), "UTF-8")

    def test_a_corrupt_member_is_exit_2_not_a_brand_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            body = ("<a:t>" + "x" * 5000 + "</a:t>").encode()
            path = self.malformed(tmp, "crc.pptx", body)
            data = bytearray(path.read_bytes())
            with zipfile.ZipFile(path) as z:
                info = z.getinfo("ppt/slides/slide1.xml")
            start = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra)
            for k in range(start + 2, start + 12):
                data[k] ^= 0xFF
            path.write_bytes(bytes(data))
            self.assert_refused(path, "crc.pptx.*(damaged|corrupt|cannot be read)")

    def test_size_and_ratio_budgets(self):
        with tempfile.TemporaryDirectory() as tmp:
            bomb = (f'<p:sld xmlns:p="{NS_P}"><!--' + " " * 3_000_000 + "--></p:sld>").encode()
            self.assert_refused(self.malformed(tmp, "ratio.pptx", bomb), "compress")
            spec = importlib.util.spec_from_file_location("check_brand_mod", CHECK)
            mod = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(mod)
            self.assertLessEqual(mod.MAX_PART, 16 * 1024 * 1024)
            mod.MAX_PART = 2000                                        # a slide master is larger
            with self.assertRaisesRegex(mod.CheckError, "too large"):
                mod.check(GOLDEN)
            mod.MAX_PART, mod.MAX_TOTAL = 16 * 1024 * 1024, 20_000      # the whole file over budget
            with self.assertRaisesRegex(mod.CheckError, "budget"):
                mod.check(GOLDEN)
            mod.MAX_TOTAL, mod.MAX_SLIDES = 256 * 1024 * 1024, 3
            with self.assertRaisesRegex(mod.CheckError, "slides"):
                mod.check(GOLDEN)

    def test_an_unexpected_error_is_exit_2_with_a_plain_message(self):
        spec = importlib.util.spec_from_file_location("check_brand_mod2", CHECK)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)

        def boom(*_):
            raise ValueError("unexpected")
        mod.check_pptx = boom
        import contextlib
        import io
        err = io.StringIO()
        with contextlib.redirect_stderr(err), contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(mod.main([str(GOLDEN)]), 2)
        self.assertIn("could not be checked", err.getvalue())

    @unittest.skipUnless(Path("/usr/bin/python3").is_file(), "no system python3")
    def test_runs_on_the_system_python_too(self):
        code, out, err = run_check(GOLDEN, "--json", python="/usr/bin/python3")
        self.assertEqual(code, 1, err)
        self.assertTrue(json.loads(out)["findings"])


def template_layout_names(template):
    """Layout names of a template, in master order (stdlib)."""
    with zipfile.ZipFile(template) as z:
        rels = z.read("ppt/slideMasters/_rels/slideMaster1.xml.rels").decode("utf-8")
        master = z.read("ppt/slideMasters/slideMaster1.xml").decode("utf-8")
        targets = {}
        for rel in re.findall(r"<Relationship [^>]*/>", rels):
            m = re.search(r'Target="\.\./slideLayouts/(slideLayout\d+\.xml)"', rel)
            if m:
                targets[re.search(r'Id="(\w+)"', rel).group(1)] = m.group(1)
        order = re.findall(r'<p:sldLayoutId [^>]*r:id="(\w+)"', master)
        return [ET.fromstring(z.read("ppt/slideLayouts/" + targets[r])).find(f"{{{NS_P}}}cSld").get("name").strip()
                for r in order]


def renamed_layout(src, dest, old, new):
    """A copy of deck `src` whose layout `old` is called `new` (a foreign layout)."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            data = zin.read(n)
            if n == layout_file(src, old):
                data = data.replace(f'<p:cSld name="{old}"'.encode(), f'<p:cSld name="{new}"'.encode())
            zout.writestr(n, data)
    return Path(dest)


def chart_xml(body):
    return ('<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
            f'xmlns:a="{NS_A}"><c:chart><c:plotArea>{body}</c:plotArea></c:chart></c:chartSpace>')


def bracket(*middle):
    """Title slide + `middle` slides + closing slide: the deck order the check wants."""
    return ([("Standard TITLE", sp("title", None, "Remote digital towers cut costs"), None)] + list(middle)
            + [("Closing Slide", "", None)])


PORTED = {"deck.first-slide", "deck.last-slide", "layout.not-company", "template.not-company",
          "footer.placeholder", "text.leftover", "shape.rounded", "chart.gridlines", "colour.chart-only",
          "text.label-headline", "slide.words"}


class TestPortedAuditRules(unittest.TestCase):
    """The owner skill's audit rules, ported into check_brand.py (design §8.5): each one fails
    a small deck built with zipfile; a deck on the full master built the right way has none."""

    def found(self, deck, year="2024"):
        return findings(deck, "--year", year)[1]

    def deck(self, name, slides, **kw):
        return build_pptx(Path(self.tmp.name) / name, slides, **kw)

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def test_check_deck_order(self):
        found = self.found(self.deck("order.pptx", [
            ("Headline (standard)", sp("title", None, "Costs fall by a third"), None)]))
        self.assertIn("deck.first-slide", rules(found, "WARN", "file"))
        self.assertIn("deck.last-slide", rules(found, "FAIL", "file"))
        found = self.found(self.deck("ok.pptx", bracket()))
        self.assertFalse(rules(found) & {"deck.first-slide", "deck.last-slide"})

    def test_check_a_layout_not_from_the_master(self):
        deck = self.deck("lay.pptx", bracket(("Headline (standard)", sp("title", None, "Costs fall by a third"), None)))
        foreign = renamed_layout(deck, Path(self.tmp.name) / "foreign.pptx", "Headline (standard)", "Title Only")
        self.assertIn("layout.not-company", rules(self.found(foreign), "FAIL", "slide", "slide 2"))
        self.assertNotIn("layout.not-company", rules(self.found(deck)))

    def test_check_a_deck_not_on_the_company_master(self):
        self.assertIn("template.not-company", rules(findings(GOLDEN)[1], "FAIL", "file"))

    def test_check_the_master_footer_placeholders(self):
        found = self.found(self.deck("footer.pptx", bracket(), footer=False))
        self.assertIn("footer.placeholder", rules(found, "FAIL", "file"))
        msg = next(f["message"] for f in found if f["rule"] == "footer.placeholder")
        self.assertIn("Presentation title", msg)
        self.assertIn("<by Presenter>", msg)

    def test_check_template_leftovers(self):
        for text in ("Lorem ipsum dolor sit amet", "Click to add text", "xxx", "Annotations: delete me"):
            with self.subTest(text=text):
                found = self.found(self.deck("left.pptx", bracket(
                    ("Headline (standard)", sp("title", None, "Costs fall by a third") + sp(None, None, text), None))))
                self.assertIn("text.leftover", rules(found, "FAIL", "slide", "slide 2"))

    def test_check_ampersand_and_exclamation_marks_stay_0_8_0_rules(self):
        found = self.found(self.deck("amp.pptx", bracket(
            ("Headline (standard)", sp("title", None, "Costs fall by a third") + sp(None, None, "Plan &amp; build now!!"), None))))
        self.assertIn("text.ampersand", rules(found, "INFO", "slide", "slide 2"))
        self.assertIn("text.exclamation", rules(found, "WARN", "slide", "slide 2"))

    def test_check_rounded_rectangles(self):
        rounded = ('<p:sp><p:nvSpPr><p:cNvPr id="50" name="Box"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
                   '<a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="004182"/>'
                   '</a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>')
        found = self.found(self.deck("round.pptx", bracket(
            ("Headline (standard)", sp("title", None, "Costs fall by a third") + rounded, None))))
        self.assertIn("shape.rounded", rules(found, "WARN", "slide", "slide 2"))

    def test_check_charts_gridlines_3d_off_palette_and_the_track_grey(self):
        grid = chart_xml('<c:barChart/><c:valAx><c:majorGridlines/></c:valAx>')
        threed = chart_xml('<c:bar3DChart/>')
        off = chart_xml('<c:barChart><c:ser><c:spPr><a:solidFill><a:srgbClr val="4472C4"/></a:solidFill>'
                        '</c:spPr></c:ser></c:barChart>')
        track = chart_xml('<c:doughnutChart><c:ser><c:dPt><c:spPr><a:solidFill><a:srgbClr val="EDF1F2"/>'
                          '</a:solidFill></c:spPr></c:dPt></c:ser></c:doughnutChart>')
        title = sp("title", None, "Load peaks in the third quarter")
        found = self.found(self.deck("charts.pptx", bracket(
            ("Headline (standard)", title, {"chart1.xml": grid}),
            ("Headline (standard)", title, {"chart2.xml": threed}),
            ("Headline (standard)", title, {"chart3.xml": off}),
            ("Headline (standard)", title, {"chart4.xml": track}),
            ("Headline (standard)", title + sp(None, None, "Track", sppr='<a:solidFill><a:srgbClr val="EDF1F2"/></a:solidFill>'), None))))
        self.assertIn("chart.gridlines", rules(found, "WARN", "slide", "slide 2"))
        self.assertIn("effect.3d", rules(found, "FAIL", "slide", "slide 3"))
        self.assertIn("colour.off-palette", rules(found, "FAIL", "slide", "slide 4"))
        self.assertEqual(rules(found, None, "slide", "slide 5"), set(), "the track grey is fine in a chart (C12)")
        self.assertIn("colour.chart-only", rules(found, "WARN", "slide", "slide 6"))
        self.assertNotIn("colour.off-palette", rules(found, None, "slide", "slide 6"))

    def test_check_a_label_headline_is_a_warning(self):
        found = self.found(self.deck("label.pptx", bracket(
            ("Headline (standard)", sp("title", None, "Next steps"), None),
            ("Divider blue world", sp("title", None, "Costs"), None))))
        self.assertIn("text.label-headline", rules(found, "WARN", "slide", "slide 2"))
        self.assertNotIn("text.label-headline", rules(found, None, "slide", "slide 3"), "dividers may be short")

    def test_check_a_wordy_slide(self):
        text = "\\n".join(["Controllers work from one remote centre for several airports"] * 12)
        found = self.found(self.deck("words.pptx", bracket(
            ("Headline (standard)", sp("title", None, "Costs fall by a third") + sp(None, None, text), None))))
        self.assertIn("slide.words", rules(found, "WARN", "slide", "slide 2"))

    def test_a_deck_on_the_full_master_built_the_right_way_has_none_of_them(self):
        deck = self.deck("full.pptx", bracket(
            ("2_Agenda", sp("title", None, "Agenda") + sp(None, 10, "Why now\\nHow it works\\nNext steps"), None),
            ("Divider blue world", sp("title", None, "Why remote towers now"), None),
            ("Sub-headline + 50:50", sp("title", None, "Two options lead to one decision")
             + sp("body", 14, "Both meet the safety case") + sp(None, 15, "Option A") + sp(None, 16, "Option B"), None),
            ("World Map | EMEA", sp("title", None, "We serve air navigation in most of EMEA"), None),
            ("Q&A", sp("title", None, "Your turn to ask"), None)), base=FULL)   # names with "&" too
        code, found = findings(deck, "--year", "2024")
        self.assertEqual(code, 0, [f for f in found if f["severity"] == "FAIL"])
        self.assertFalse(rules(found) & PORTED, [f for f in found if f["rule"] in PORTED])

    def test_the_tokens_list_the_44_layouts_of_the_two_templates(self):
        tokens = json.loads(TOKENS.read_text(encoding="utf-8"))
        full, slim = template_layout_names(FULL), set(template_layout_names(TEMPLATE))
        self.assertEqual(len(full), 44)
        self.assertEqual(len(slim), 25)
        self.assertEqual([(l["name"], l["templates"]) for l in tokens["pptx"]["layouts"]],
                         [(n, "slim and full" if n in slim else "full only") for n in full])

    def test_frq_pptx_audit_only_delegates_to_check_brand(self):
        tree = ast.parse(FRQ_PPTX.read_text(encoding="utf-8"), feature_version=(3, 9))
        funcs = {n.name: n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef)}
        self.assertNotIn("audit", funcs, "no second set of rules")
        body = ast.dump(funcs["cmd_audit"])
        self.assertIn("'check_brand'", body)
        self.assertIn("attr='main'", body)
        self.assertLess(len(funcs["cmd_audit"].body), 6)
        src = FRQ_PPTX.read_text(encoding="utf-8")
        for rule_words in ("LEFTOVER", "US_WORDS", "CONTRACTIONS", "GRADIENT_OK", "def _is_title_case"):
            self.assertNotIn(rule_words, src)

    def test_the_audit_command_prints_the_check_brand_result(self):
        r = subprocess.run([sys.executable, "-I", str(FRQ_PPTX), "audit", str(GOLDEN), "--json"],
                           capture_output=True, text=True, check=False)
        self.assertEqual(r.returncode, 1, r.stderr)
        code, found = findings(GOLDEN)
        self.assertEqual([f["rule"] for f in json.loads(r.stdout)["findings"]], [f["rule"] for f in found])

    def test_no_bare_pip_install_in_the_skill(self):
        for rel in skill_files():
            if not is_text(rel):
                continue
            for line in (SKILL / rel).read_text(encoding="utf-8").splitlines():
                if re.search(r"\\bpip3? install\\b", line):
                    self.assertIn("~/.ai-sdlc/venv", line, rel)

    def test_new_deck_builds_through_frq_pptx_on_the_slim_template(self):
        tree = ast.parse(NEW_DECK.read_text(encoding="utf-8"), feature_version=(3, 9))
        calls = [ast.dump(n.func) for n in ast.walk(tree) if isinstance(n, ast.Call)]
        self.assertTrue([c for c in calls if "id='frq_pptx'" in c and "attr='build'" in c])
        self.assertNotIn("Presentation(", NEW_DECK.read_text(encoding="utf-8"), "one builder")
        self.assertEqual(load_new_deck().TEMPLATE, "slim")


def load_assets_from(path):
    spec = importlib.util.spec_from_file_location("brand_assets_copy", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class TestAssetResolver(unittest.TestCase):
    """scripts/brand_assets.py: the placed folder first, then the kit copy (design §8.7)."""

    def test_assets_are_found_in_the_placed_folder_then_the_kit_copy(self):
        rel = "templates/frq-master.pptx"
        kit_rel = ".ai-sdlc/kit/template/.claude/skills/frq-brandbook"
        with tempfile.TemporaryDirectory() as tmp:
            root = helpers.make_repo(Path(tmp).resolve() / "repo", {"README.md": "team\\n"})
            placed = root / PLACED
            (placed / "scripts").mkdir(parents=True)
            (placed / "scripts/brand_assets.py").write_bytes(ASSETS.read_bytes())
            sub = root / "docs/decks"
            sub.mkdir(parents=True)
            mod = load_assets_from(placed / "scripts/brand_assets.py")
            with self.assertRaises(mod.AssetMissing) as cm:
                mod.asset_path(rel, start=sub)
            self.assertIn(f"{kit_rel}/assets/{rel}", str(cm.exception))
            self.assertIn("check the kit", str(cm.exception))
            kit = root / kit_rel / "assets/templates"
            kit.mkdir(parents=True)
            (kit / "frq-master.pptx").write_bytes(b"kit copy")
            self.assertEqual(mod.asset_path(rel, start=sub), kit / "frq-master.pptx")
            self.assertEqual(mod.asset_path(rel), kit / "frq-master.pptx", "from the script's folder too")
            (placed / "assets/templates").mkdir(parents=True)
            (placed / "assets/templates/frq-master.pptx").write_bytes(b"placed")
            self.assertEqual(mod.asset_path(rel, start=sub), placed / "assets/templates/frq-master.pptx")
            with self.assertRaises(mod.AssetMissing):
                mod.find("../../outside.txt")
            r = subprocess.run([sys.executable, "-I", str(placed / "scripts/brand_assets.py"), "keyvisual/none.jpg"],
                               capture_output=True, text=True, check=False, cwd=sub)
            self.assertEqual(r.returncode, 2)
            self.assertNotIn("Traceback", r.stderr)
            self.assertFalse(list(placed.rglob("__pycache__")))

    def test_in_the_kit_every_script_finds_its_assets(self):
        mod = load_assets_from(ASSETS)
        self.assertEqual(mod.asset_path("templates/frq-master.pptx"), FULL)
        self.assertEqual(mod.find("brand-tokens.json"), TOKENS)


def load_frq_pptx():
    spec = importlib.util.spec_from_file_location("frq_pptx_mod", FRQ_PPTX)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def load_new_deck():
    spec = importlib.util.spec_from_file_location("new_deck_mod", NEW_DECK)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class TestNewDeckClassification(unittest.TestCase):
    """E2E 2026-10-09: Copilot imported new_deck and called build()/set_footer() with a class
    it picked itself. The internals check the value too; no python-pptx needed for this."""

    def test_build_and_set_footer_take_the_class_by_keyword_only_and_check_it(self):
        nd = load_new_deck()
        with tempfile.TemporaryDirectory() as tmp:
            outline = Path(tmp) / "o.md"
            outline.write_text("# T\n\n## One\n- a\n", encoding="utf-8")
            out = Path(tmp) / "d.pptx"
            with self.assertRaises(TypeError):
                nd.build(outline, out, "Frequentis General", 2026)          # positional: refused
            for guess in ("internal", "Confidential", "", None):
                with self.subTest(guess=guess), self.assertRaises(ValueError):
                    nd.build(outline, out, classification=guess)
                with self.subTest(guess=guess), self.assertRaises(ValueError):
                    nd.set_footer(object(), classification=guess, year=2026, title="T")
            self.assertFalse(out.exists())
        self.assertEqual(nd.ALLOWED, nd.CLASSES + ("Frequentis [classification to be set]",))
        for value in nd.ALLOWED:
            self.assertEqual(nd.checked_classification(value), value)


@unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx is not installed here")
class TestNewDeck(unittest.TestCase):
    def test_new_deck_from_an_outline_passes_the_check(self):
        with tempfile.TemporaryDirectory() as tmp:
            outline = Path(tmp) / "outline.md"
            outline.write_text("# Remote digital towers\nAna Pop, 9 October 2026\n\n"
                               "## Controllers see more with fewer screens\n"
                               "> What changes for the tower team\n- One working position\n"
                               "  - Shared voice and data\n- Fewer handovers\n\n"
                               "## Next we pilot in one tower\n- Pilot in one tower\n", encoding="utf-8")
            out = Path(tmp) / "deck.pptx"
            r = subprocess.run([sys.executable, str(NEW_DECK), str(outline), str(out),
                                "--classification", "Frequentis General", "--year", "2026"],
                               capture_output=True, text=True, check=False)
            self.assertEqual(r.returncode, 0, r.stderr)
            from pptx import Presentation
            prs = Presentation(str(out))
            self.assertEqual([s.slide_layout.name for s in prs.slides],
                             ["Standard TITLE", "Sub-headline + field", "Headline + field", "Closing Slide"])
            code, found = findings(out, "--year", "2026")
            self.assertEqual([f for f in found if f["severity"] != "INFO"], [])
            self.assertEqual(code, 0)
            r = subprocess.run([sys.executable, str(NEW_DECK), str(outline), str(out),
                                "--classification", "Frequentis General"],
                               capture_output=True, text=True, check=False)
            self.assertNotEqual(r.returncode, 0, "never overwrites")

    def test_the_shape_table_and_chart_recipe_in_building_decks_is_on_brand(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        doc = (SKILL / "references/building-decks.md").read_text(encoding="utf-8")
        recipe = doc.split("### Shapes, free text, tables and charts", 1)[1].split("```python", 1)[1].split("```", 1)[0]
        ns = {}
        exec(compile(recipe, "building-decks.md", "exec"), ns)
        Inches = ns["Inches"]
        with tempfile.TemporaryDirectory() as tmp:
            prs = Presentation(str(TEMPLATE))
            layouts = {l.name: l for l in prs.slide_layouts}
            s = prs.slides.add_slide(layouts["Standard TITLE"])     # a deck starts on the title slide
            s.shapes.title.text = "Remote digital towers"
            fp = load_frq_pptx()
            fp.set_footer(prs, classification="Frequentis General", year=2024, title="Recipes",
                          presenter="by Ana Pop")
            s = prs.slides.add_slide(layouts["Headline (standard)"])
            s.shapes.title.text = "Three steps to a decision"
            ns["flat_box"](s, Inches(0.5), Inches(1.2), Inches(2), Inches(1))
            ns["text_box"](s, Inches(3), Inches(1.2), Inches(4), Inches(1), "Callout text")
            s = prs.slides.add_slide(layouts["Headline (standard)"])
            s.shapes.title.text = "Option B costs less"
            ns["brand_table"](s, [("Option", "Cost"), ("A", "2,115"), ("B", "1,500")],
                              Inches(0.5), Inches(1.2), Inches(5))
            s = prs.slides.add_slide(layouts["Headline (standard)"])
            s.shapes.title.text = "Load peaks in the third quarter"
            data = CategoryChartData()
            data.categories = ["Q1", "Q2", "Q3"]
            data.add_series("Load", (3, 4, 7))
            chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.2),
                                       Inches(8), Inches(3.5), data).chart
            ns["brand_chart"](chart)
            prs.slides.add_slide(layouts["Closing Slide"])           # and ends on the closing slide
            out = Path(tmp) / "recipe.pptx"
            prs.save(str(out))
            code, found = findings(out, "--year", "2024")
            self.assertEqual([f for f in found if f["severity"] != "INFO"], [])

    def run_new_deck(self, tmp, text, *args):
        outline = Path(tmp) / "o.md"
        outline.write_text(text, encoding="utf-8")
        out = Path(tmp) / f"d{len(list(Path(tmp).iterdir()))}.pptx"
        r = subprocess.run([sys.executable, str(NEW_DECK), str(outline), str(out), *args],
                           capture_output=True, text=True, check=False)
        return r, out

    def test_the_classification_is_never_guessed(self):
        with tempfile.TemporaryDirectory() as tmp:
            r, out = self.run_new_deck(tmp, "# Title\n\n## One\n- a\n")
            self.assertNotEqual(r.returncode, 0)
            self.assertIn("classification", r.stderr)
            self.assertFalse(out.exists())

    def test_the_placeholder_builds_and_the_check_flags_it(self):
        with tempfile.TemporaryDirectory() as tmp:
            r, out = self.run_new_deck(tmp, "# Title\n\n## One\n- a\n",
                                       "--classification", "Frequentis [classification to be set]")
            self.assertEqual(r.returncode, 0, r.stderr)
            self.assertIn("classification to be set", r.stderr)
            code, found = findings(out)
            self.assertEqual(code, 0)
            self.assertIn("footer.classification-to-set", rules(found, "WARN"))
            r, out = self.run_new_deck(tmp, "# Title\n\n## One\n- a\n", "--classification", "Internal")
            self.assertNotEqual(r.returncode, 0)
            self.assertFalse(out.exists())

    def test_numbered_lists_paragraphs_and_bold_are_kept(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as tmp:
            r, out = self.run_new_deck(tmp, "# Risks\n\n## Two risks need an owner\n"
                                       "1. **Alarm routing** is late\n2) Spare parts\n"
                                       "Both need a decision this sprint.\n\n## Divider only\n",
                                       "--classification", "Frequentis Confidential")
            self.assertEqual(r.returncode, 0, r.stderr)
            slides = list(Presentation(str(out)).slides)
            body = " / ".join(sh.text_frame.text for sh in slides[1].placeholders
                              if sh.placeholder_format.idx == 15)
            self.assertIn("Alarm routing is late", body)
            self.assertIn("Spare parts", body)
            self.assertIn("Both need a decision this sprint.", body)
            self.assertNotIn("**", body)
            self.assertEqual(slides[2].slide_layout.name, "Headline (standard)")

    def test_lines_it_cannot_place_are_reported_and_an_empty_slide_stops_it(self):
        with tempfile.TemporaryDirectory() as tmp:
            r, out = self.run_new_deck(tmp, "# T\n\n## Costs\n| a | b |\n|---|---|\n",
                                       "--classification", "Frequentis General")
            self.assertNotEqual(r.returncode, 0)
            self.assertIn("| a | b |", r.stderr)
            self.assertFalse(out.exists())
            r, out = self.run_new_deck(tmp, "# T\n\n## Costs\n- one\n```\ncode\n```\n",
                                       "--classification", "Frequentis General")
            self.assertEqual(r.returncode, 0, r.stderr)
            self.assertIn("not placed", r.stderr)


MANY_LAYOUTS = {"title": "Remote digital towers", "presenter": "by Ana Pop", "slides": [
    {"layout": "Standard TITLE", "title": "Remote digital towers cut staffing costs",
     "subtitle": "Customer briefing, 9 October 2026"},
    {"layout": "2_Agenda", "title": "Agenda",
     "content": [["Why remote towers now", "How the solution works", "Business case"]]},
    {"layout": "Divider blue world", "title": "Why remote towers now"},
    {"layout": "Sub-headline + 50:50", "title": "Two options lead to one decision",
     "subtitle": "Both meet the safety case",
     "content": [["Upgrade the tower in place"], {"table": {"header": ["Option", "Cost"],
                                                            "rows": [["A", "2,115"], ["B", "1,500"]]}}]},
    {"layout": "World Map | EMEA", "title": "We serve air navigation across EMEA"},
    {"layout": "Headline (standard)", "title": "Delivery runs in five phases"},
    {"layout": "Closing Slide"}]}
SLIM_ONLY = {"title": "Status", "slides": [
    {"layout": "Standard TITLE", "title": "The pilot tower is ready for the trial"},
    {"layout": "Headline + field", "title": "Two risks need an owner this sprint",
     "content": [["Alarm routing is late", "Spare parts arrive in week three"]]},
    {"layout": "Closing Slide"}]}


def run_frq_pptx(*args, cwd=None, python=None, flags=("-I",), env=None):
    r = subprocess.run([python or sys.executable, *flags, str(FRQ_PPTX), *map(str, args)],
                       capture_output=True, text=True, check=False, cwd=cwd, env=env)
    return r


def count_layouts(deck):
    with zipfile.ZipFile(deck) as z:
        return len([n for n in z.namelist() if re.fullmatch(r"ppt/slideLayouts/slideLayout\d+\.xml", n)])


class TestBuilderWithoutPptx(unittest.TestCase):
    """frq_pptx.py guarantees that hold without python-pptx (design §8.5)."""

    def test_without_python_pptx_it_says_how_to_get_it(self):
        with tempfile.TemporaryDirectory() as tmp:
            spec = Path(tmp) / "spec.json"
            spec.write_text(json.dumps(SLIM_ONLY), encoding="utf-8")
            r = run_frq_pptx("build", spec, Path(tmp) / "out.pptx", "--classification", "Frequentis General",
                             flags=("-I", "-S"))                        # -S: no site-packages, no pptx
            self.assertEqual(r.returncode, 2, r.stdout + r.stderr)
            self.assertNotIn("Traceback", r.stderr)
            self.assertIn("~/.ai-sdlc/venv", r.stderr)
            self.assertIn("ai-sdlc-doc-powerpoint", r.stderr)
            self.assertFalse((Path(tmp) / "out.pptx").exists())

    def test_the_classification_is_required_and_checked_before_anything_else(self):
        fp = load_frq_pptx()
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "d.pptx"
            with self.assertRaises(TypeError):
                fp.build(SLIM_ONLY, out)                                 # no classification= : refused
            for guess in ("secret", "General", "internal", "", None):
                with self.subTest(guess=guess), self.assertRaises(ValueError):
                    fp.build(SLIM_ONLY, out, classification=guess)
                with self.subTest(guess=guess), self.assertRaises(ValueError):
                    fp.set_footer(object(), classification=guess, year=2026)
            with self.assertRaises(TypeError):
                fp.set_footer(object(), "Frequentis General", 2026)      # positional: refused
            with self.assertRaises(ValueError):                         # a spec cannot pick another class
                fp.build(dict(SLIM_ONLY, classification="Frequentis Public"), out,
                         classification="Frequentis General")
            self.assertFalse(out.exists())
            r = run_frq_pptx("build", Path(tmp) / "s.json", out)
            self.assertEqual(r.returncode, 2)
            self.assertIn("--classification", r.stderr)
        self.assertEqual(fp.ALLOWED, fp.CLASSES + ("Frequentis [classification to be set]",))
        self.assertEqual(fp.SERIES[:2], ["004182", "00AAE1"], "the palette comes from brand-tokens.json")

    def test_render_writes_only_under_ai_sdlc_tmp(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = helpers.make_repo(Path(tmp).resolve() / "repo", {"README.md": "team\n"})
            bin_ = Path(tmp) / "bin"
            bin_.mkdir()
            log = Path(tmp) / "soffice.log"
            (bin_ / "soffice").write_text(
                "#!/bin/sh\n"
                f'echo "$@" >> "{log}"\n'
                'out=""; deck=""\n'
                'while [ $# -gt 0 ]; do case "$1" in --outdir) out="$2"; shift;; *.pptx) deck="$1";; esac; shift; done\n'
                'name=$(basename "$deck" .pptx); echo pdf > "$out/$name.pdf"\n')
            (bin_ / "pdftoppm").write_text('#!/bin/sh\nfor last; do :; done\necho png > "$last-1.png"\n')
            for f in bin_.iterdir():
                f.chmod(0o755)
            deck = root / "docs/deck.pptx"
            deck.parent.mkdir()
            deck.write_bytes(GOLDEN.read_bytes())
            env = dict(os.environ, PATH=f"{bin_}{os.pathsep}{os.environ.get('PATH', '')}")
            r = run_frq_pptx("render", deck, cwd=root, env=env)
            self.assertEqual(r.returncode, 0, r.stderr)
            out = root / ".ai-sdlc/tmp/deck"
            self.assertTrue((out / "deck.pdf").is_file())
            self.assertTrue((out / "slide-1.png").is_file())
            self.assertIn((root / ".ai-sdlc/tmp/lo-profile").as_uri(), log.read_text())
            r = run_frq_pptx("render", deck, Path(tmp) / "elsewhere", cwd=root, env=env)
            self.assertEqual(r.returncode, 2)
            self.assertIn(".ai-sdlc", r.stderr)
            self.assertFalse((Path(tmp) / "elsewhere").exists())
            self.assertEqual(deck.read_bytes(), GOLDEN.read_bytes(), "the deck is never changed")
            r = run_frq_pptx("render", deck, cwd=root, env=dict(os.environ, PATH=str(Path(tmp) / "none")))
            self.assertEqual(r.returncode, 2)
            self.assertIn("PowerPoint", r.stderr)


@unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx is not installed here")
class TestBuilder(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)

    def spec(self, data, name="spec.json"):
        p = self.dir / name
        p.write_text(json.dumps(data), encoding="utf-8")
        return p

    def test_build_from_a_spec_with_many_layouts_passes_the_check(self):
        out = self.dir / "deck.pptx"
        r = run_frq_pptx("build", self.spec(MANY_LAYOUTS), out, "--classification", "Frequentis General",
                         "--year", "2026")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("used the full master for: World Map | EMEA", r.stdout)
        # the timeline recipe from references/build-spec.md: chevrons and an axis, one line weight
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt
        F = load_frq_pptx()
        prs = Presentation(str(out))
        slide = prs.slides[5]
        for i, name in enumerate(["Analysis", "Design", "Build", "Test", "Operate"]):
            c = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(0.47 + i * 1.8), Inches(1.2), Inches(1.75), Inches(0.5))
            c.fill.solid()
            c.fill.fore_color.rgb = F.rgb(F.LIGHT_BLUE if i == 1 else F.WARM_GREY)
            F.flat(c)
        ln = slide.shapes.add_connector(1, Inches(0.47), Inches(3), Inches(9.53), Inches(3))
        ln.line.color.rgb = F.rgb(F.BLUE)
        ln.line.width = Pt(1.5)
        ln.shadow.inherit = False
        final = self.dir / "deck-timeline.pptx"
        prs.save(str(final))
        self.assertEqual([s.slide_layout.name.strip() for s in prs.slides],
                         [s["layout"] for s in MANY_LAYOUTS["slides"]])
        code, found = findings(final, "--year", "2026")
        self.assertEqual(code, 0, [f for f in found if f["severity"] == "FAIL"])
        self.assertFalse(rules(found) & PORTED, [f for f in found if f["rule"] in PORTED])
        footer = " | ".join(t for t in re.findall(r"<a:t>([^<]*)</a:t>", zipfile.ZipFile(final).read(
            "ppt/slideMasters/slideMaster1.xml").decode("utf-8")) if "Frequentis" in t or "Ana" in t or "towers" in t)
        for part in ("Remote digital towers", "by Ana Pop", "Frequentis General", "© Frequentis AG 2026"):
            self.assertIn(part, footer)
        r = run_frq_pptx("audit", final)                          # the audit command is the check
        self.assertEqual(r.returncode, 0, r.stdout)
        self.assertIn("Brand check:", r.stdout)

    # --- re-test round 2 (2026-10-10): the ATM key visual on the title slide by default --

    def title_picture(self, deck):
        """The bytes of the big picture on the title slide's layout (the key visual square)."""
        from pptx import Presentation
        prs = Presentation(str(deck))
        lay = prs.slides[0].slide_layout
        shape = load_frq_pptx().key_visual_shape(lay)
        rid = shape._element.find(".//" + "{http://schemas.openxmlformats.org/drawingml/2006/main}blip").get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        return lay.part.related_part(rid).blob

    def kv(self, bu):
        tokens = json.loads((SKILL / "brand-tokens.json").read_text(encoding="utf-8"))
        rel = next(b["key_visual"] for b in tokens["business_units"] if b["id"] == bu)
        return (SKILL / rel).read_bytes()

    def test_the_title_slide_shows_the_atm_key_visual_by_default(self):
        fp = load_frq_pptx()
        template = SKILL / "assets/templates/frq-template-slim-core.pptx"
        before = hashlib.sha256(template.read_bytes()).hexdigest()
        out = self.dir / "atm.pptx"
        fp.build(SLIM_ONLY, out, classification="Frequentis General", year=2026)
        self.assertEqual(self.title_picture(out), self.kv("ATM"))
        self.assertEqual(hashlib.sha256(template.read_bytes()).hexdigest(), before, "template untouched")
        code, found = findings(out, "--year", "2026")
        self.assertEqual(code, 0, [f for f in found if f["severity"] == "FAIL"])

    def test_another_business_unit_or_none_on_request(self):
        fp = load_frq_pptx()
        mar = self.dir / "mar.pptx"
        fp.build(dict(SLIM_ONLY, business_unit="MAR"), mar, classification="Frequentis General", year=2026)
        self.assertEqual(self.title_picture(mar), self.kv("MAR"))
        plain = self.dir / "plain.pptx"
        fp.build(dict(SLIM_ONLY, key_visual=False), plain, classification="Frequentis General", year=2026)
        globe = self.dir / "globe.pptx"
        fp.build(dict(SLIM_ONLY, business_unit="CORP"), globe, classification="Frequentis General", year=2026)
        self.assertEqual(self.title_picture(plain), self.title_picture(globe))
        self.assertNotEqual(self.title_picture(plain), self.kv("ATM"))
        r = run_frq_pptx("build", self.spec(SLIM_ONLY), self.dir / "def.pptx", "--classification",
                         "Frequentis General", "--business-unit", "DEF")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(self.title_picture(self.dir / "def.pptx"), self.kv("DEF"))
        with self.assertRaises(ValueError):
            fp.build(dict(SLIM_ONLY, business_unit="XYZ"), self.dir / "x.pptx",
                     classification="Frequentis General", year=2026)

    def test_new_deck_uses_the_atm_key_visual_too(self):
        outline = self.dir / "o.md"
        outline.write_text("# Remote towers\n\n## The pilot works\n- One tower\n", encoding="utf-8")
        for args, bu in (([], "ATM"), (["--business-unit", "PS"], "PS")):
            out = self.dir / f"nd-{bu}.pptx"
            r = subprocess.run([sys.executable, str(NEW_DECK), str(outline), str(out),
                                "--classification", "Frequentis General", *args],
                               capture_output=True, text=True, check=False)
            self.assertEqual(r.returncode, 0, r.stderr)
            self.assertEqual(self.title_picture(out), self.kv(bu))

    # --- re-test round 2, N4: a map pin from latitude and longitude -------------------------

    CITIES = {   # city: (lat, lon, the country shape it must fall in)
        "Vienna": (48.21, 16.37, "Austria"), "Madrid": (40.42, -3.70, "Spain"),
        "Sydney": (-33.87, 151.21, "Australia"), "Cape Town": (-33.92, 18.42, "South Africa"),
        "Reykjavik": (64.15, -21.94, "Iceland"), "Nairobi": (-1.29, 36.82, "Kenya"),
        "Tokyo": (35.68, 139.69, "Japan"), "Sao Paulo": (-23.55, -46.63, "Brazil"),
        "Chicago": (41.88, -87.63, "USA"), "Dubai": (25.20, 55.27, "United Arab Emirates"),
        "Warsaw": (52.23, 21.01, "Poland"), "Tehran": (35.69, 51.39, "Iran"),
    }

    def country_box(self, layout, name):
        """The country shape's box on the slide, in inches (through the group transforms)."""
        A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

        def tf_of(el, outer):
            xf = el.find(A + "xfrm") if el.find(A + "xfrm") is not None else el.find(".//" + A + "xfrm")
            o, e, co, ce = (xf.find(A + k) for k in ("off", "ext", "chOff", "chExt"))
            sx, sy = int(e.get("cx")) / int(ce.get("cx")), int(e.get("cy")) / int(ce.get("cy"))
            return lambda x, y: outer(int(o.get("x")) + (x - int(co.get("x"))) * sx,
                                      int(o.get("y")) + (y - int(co.get("y"))) * sy)

        def walk(shapes, tf):
            for sh in shapes:
                if sh.shape_type == 6:
                    found = walk(sh.shapes, tf_of(sh._element.find(".//{http://schemas.openxmlformats.org/"
                                                                   "presentationml/2006/main}grpSpPr"), tf))
                    if found:
                        return found
                elif sh.name == name:
                    x0, y0 = tf(sh.left, sh.top)
                    x1, y1 = tf(sh.left + sh.width, sh.top + sh.height)
                    return [v / 914400 for v in (x0, y0, x1, y1)]
            return None
        group = max((sh for sh in layout.shapes if sh.shape_type == 6), key=lambda g: len(g.shapes))
        return walk(group.shapes, tf_of(group._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}grpSpPr"), lambda x, y: (x, y)))

    def test_map_points_land_in_their_country_on_the_world_and_emea_maps(self):
        from pptx import Presentation
        fp = load_frq_pptx()
        prs = Presentation(str(FULL))
        for lname in ("World Map", "Map", "World Map | Americas", "1_MAP APAC", "World Map | EMEA"):
            lay = next(l for l in prs.slide_layouts if l.name.strip() == lname)
            for city, (lat, lon, country) in self.CITIES.items():
                with self.subTest(layout=lname, city=city):
                    try:
                        x, y = fp.map_point(lname, lat, lon)
                    except ValueError:
                        self.assertEqual(lname, "World Map | EMEA", city)
                        self.assertIn(city, ("Sydney", "Tokyo", "Sao Paulo", "Chicago", "Reykjavik"))
                        continue
                    x0, y0, x1, y1 = self.country_box(lay, country)
                    self.assertTrue(x0 <= x <= x1 and y0 <= y <= y1, (x, y, (x0, y0, x1, y1)))
        with self.assertRaises(ValueError):
            fp.map_point("Headline + field", 48.2, 16.4)          # not a calibrated map

    def test_pins_in_a_spec_are_placed_and_pass_the_check(self):
        from pptx import Presentation
        fp = load_frq_pptx()
        spec = {"title": "Sites", "slides": [
            {"layout": "Standard TITLE", "title": "Our towers in Europe and Africa"},
            {"layout": "World Map | EMEA", "title": "Three sites run the new tower",
             "pins": [{"lat": 48.21, "lon": 16.37, "label": "Vienna"},
                      {"lat": -1.29, "lon": 36.82, "label": "Nairobi"}, {"lat": 40.42, "lon": -3.70}]},
            {"layout": "Closing Slide"}]}
        out = self.dir / "pins.pptx"
        fp.build(spec, out, classification="Frequentis General", year=2026)
        slide = Presentation(str(out)).slides[1]
        pins = [sh for sh in slide.shapes if sh.name.startswith("Map pin")]
        self.assertEqual(len(pins), 3)
        x, y = fp.map_point("World Map | EMEA", 48.21, 16.37)
        vienna = next(sh for sh in pins if sh.name == "Map pin: Vienna")
        cx, cy = (vienna.left + vienna.width / 2) / 914400, (vienna.top + vienna.height / 2) / 914400
        self.assertAlmostEqual(cx, x, places=2)
        self.assertAlmostEqual(cy, y, places=2)
        self.assertIn("Vienna", [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame])
        code, found = findings(out, "--year", "2026")
        self.assertEqual(code, 0, [f for f in found if f["severity"] == "FAIL"])

    def test_build_and_footer_need_a_checked_classification_by_keyword(self):
        fp = load_frq_pptx()
        out = self.dir / "p.pptx"
        result = fp.build(SLIM_ONLY, out, classification="Frequentis [classification to be set]", year=2026)
        self.assertEqual(result["template"], "slim")
        code, found = findings(out, "--year", "2026")
        self.assertEqual(code, 0)
        self.assertIn("footer.classification-to-set", rules(found, "WARN"))
        r = run_frq_pptx("footer", out, self.dir / "p2.pptx", "--classification", "Frequentis Confidential",
                         "--title", "Status", "--presenter", "by Ana Pop", "--year", "2026")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("Frequentis Confidential", r.stdout)
        r = run_frq_pptx("footer", out, self.dir / "p3.pptx", "--classification", "Confidential")
        self.assertEqual(r.returncode, 2)
        self.assertFalse((self.dir / "p3.pptx").exists())

    def test_it_never_overwrites_its_input_or_an_existing_file(self):
        out = self.dir / "d.pptx"
        self.assertEqual(run_frq_pptx("build", self.spec(SLIM_ONLY), out, "--classification",
                                      "Frequentis General").returncode, 0)
        before = out.read_bytes()
        for args in (("footer", out, out, "--classification", "Frequentis General"),
                     ("build", self.spec(SLIM_ONLY), out, "--classification", "Frequentis General")):
            r = run_frq_pptx(*args)
            self.assertEqual(r.returncode, 2, args)
            self.assertNotIn("Traceback", r.stderr)
            self.assertEqual(out.read_bytes(), before)

    def test_the_slim_template_is_the_default_and_the_full_master_only_when_needed(self):
        slim_out, full_out = self.dir / "slim.pptx", self.dir / "full.pptx"
        r = run_frq_pptx("build", self.spec(SLIM_ONLY), slim_out, "--classification", "Frequentis General")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(count_layouts(slim_out), 25)
        self.assertNotIn("full master", r.stdout)
        r = run_frq_pptx("build", self.spec(MANY_LAYOUTS), full_out, "--classification", "Frequentis General")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(count_layouts(full_out), 44)
        self.assertIn("used the full master for: World Map | EMEA", r.stdout)
        self.assertLess(slim_out.stat().st_size * 4, full_out.stat().st_size)
        r = run_frq_pptx("build", self.spec(MANY_LAYOUTS), self.dir / "x.pptx", "--classification",
                         "Frequentis General", "--template", "slim")
        self.assertEqual(r.returncode, 2)
        self.assertIn("World Map | EMEA", r.stderr)
        self.assertNotIn("Traceback", r.stderr)
        r = run_frq_pptx("build", self.spec(SLIM_ONLY), self.dir / "f.pptx", "--classification",
                         "Frequentis General", "--template", "full")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(count_layouts(self.dir / "f.pptx"), 44)

    def test_new_deck_output_has_the_slim_templates_25_layouts(self):
        outline = self.dir / "o.md"
        outline.write_text("# Remote digital towers\n\n## Controllers see more with fewer screens\n- One position\n",
                           encoding="utf-8")
        r = subprocess.run([sys.executable, str(NEW_DECK), str(outline), str(self.dir / "n.pptx"),
                            "--classification", "Frequentis General"], capture_output=True, text=True, check=False)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(count_layouts(self.dir / "n.pptx"), 25)


KIT_PREFIX = ".ai-sdlc/kit/template/.claude/skills/frq-brandbook/"
KIT_ONLY_MENTION = re.compile(r"[\w./-]*assets/(?:(?:templates|layouts|examples|keyvisual)/[\w./-]*"
                              r"|background/[\w.-]*\.jpg|logo/[\w.-]*\.png)")


def placed_text():
    """{rel: text} of SKILL.md and the references as setup places them."""
    files = place.placed_skill_files(KIT, "frq-brandbook")
    out = {}
    for path, data in files.items():
        rel = path.split("/ai-sdlc-frq-brandbook/", 1)[1]
        if rel == "SKILL.md" or (rel.startswith("references/") and rel.endswith(".md")):
            out[rel] = data
    return out


class TestSkillText(unittest.TestCase):
    """The merged skill text (design §8.6), references and PROVENANCE (Task B4)."""

    def setUp(self):
        self.skill = (SKILL / "SKILL.md").read_text(encoding="utf-8")

    def test_the_create_and_apply_workflows(self):
        for needed in ("message pyramid", "Executive", "Self-explanatory", "Apply", "Must", "Should",
                       "into a new file", "one question at a time", "Outline first, then stop",
                       "Build nothing until the person agrees", "all fixes, only the Must fixes, or pick by slide number",
                       "German decks", "frq_pptx.py build", "new_deck.py"):
            self.assertIn(needed, self.skill)

    def test_retest_round_2_one_question_atm_title_and_decks_folder(self):
        create = self.skill.split("## Create", 1)[1].split("## Check", 1)[0]
        apply_ = self.skill.split("## Apply an existing file", 1)[1].split("## Per deliverable", 1)[0]
        for part in (create, apply_):
            self.assertIn("Ask the classification alone first", part)
            self.assertIn("then the rest one at a time", part)
        self.assertNotIn("keep `Standard TITLE` as it is", self.skill)
        self.assertIn("the ATM key visual on the title slide", create)
        self.assertIn("--business-unit", create)
        self.assertIn("the big square shows the unit's key visual and the five small unit tiles stay", create)
        self.assertIn("`docs/decks/`", create)
        # a one-shot run asked the classification and the presenter together (re-test round 2)
        self.assertIn("Also when nobody can answer in this run, ask only the classification and stop", create)
        self.assertIn("never assume", create.split("`docs/decks/`", 1)[1][:200])

    def test_layouts_reference_names_all_44_and_links_each_preview(self):
        text = (SKILL / "references/layouts.md").read_text(encoding="utf-8")
        heads = re.findall(r"(?m)^## (\d+)\. (.+)$", text)
        full, slim = template_layout_names(FULL), set(template_layout_names(TEMPLATE))
        self.assertEqual([int(n) for n, _ in heads], list(range(1, 45)))
        self.assertEqual([name.strip() for _, name in heads], full)
        for section, (_, name) in zip(re.split(r"(?m)^## \d+\. ", text)[1:], heads):
            with self.subTest(layout=name):
                where = "slim and full" if name.strip() in slim else "full only"
                self.assertIn(f"Templates: **{where}**", section)
                link = re.search(r"\]\(\.\./(assets/layouts/[\w.-]+\.jpg)\)", section)
                self.assertTrue(link and (SKILL / link.group(1)).is_file(), name)
                self.assertIn(KIT_PREFIX + link.group(1), section)

    def test_kit_only_files_are_named_by_exact_kit_path(self):
        """Copilot's search skips .ai-sdlc/: every kit-only file is named by its full kit path.
        Placement rewrites each link to a kit-only file into .ai-sdlc/kit (Task K1); the link
        must land on a real kit-only file there."""
        for rel, text in placed_text().items():
            targets = re.findall(r"\]\(([^)\s]+)\)", text)
            for target in targets:
                if KIT_ONLY_MENTION.search(target):
                    landed = posixpath.normpath(posixpath.join(PLACED, posixpath.dirname(rel),
                                                               target.split("#")[0]))
                    self.assertTrue(landed.startswith(KIT_PREFIX), (rel, target))
                    inside = landed[len(KIT_PREFIX):]
                    self.assertTrue((SKILL / inside).is_file(), (rel, target))
                    self.assertTrue(is_kit_only(inside), (rel, target))
            bare = re.sub(r"\]\([^)\s]+\)", "]()", text)
            for m in KIT_ONLY_MENTION.finditer(bare):
                with self.subTest(file=rel, mention=m.group(0)):
                    self.assertTrue(m.group(0).startswith(KIT_PREFIX + "assets/"), m.group(0))

    def test_atm_is_the_default_business_unit(self):
        for rel in ("SKILL.md", "references/brand-rules.md"):
            text = (SKILL / rel).read_text(encoding="utf-8")
            self.assertRegex(text, r"ATM[^.\n]{0,40}by default|Default BU: ATM", rel)
            self.assertIn("names another", text, rel)
        self.assertEqual(json.loads(TOKENS.read_text(encoding="utf-8"))["default_business_unit"], "atm")

    def test_build_spec_reference_paths_exist(self):
        text = (SKILL / "references/build-spec.md").read_text(encoding="utf-8")
        found = set(re.findall(r"(?<![\w-])((?:assets|scripts)/[\w./-]*[\w])", text))
        self.assertTrue({"scripts/frq_pptx.py", "scripts/check_brand.py"} <= found)
        for rel in sorted(found):
            self.assertTrue((SKILL / rel).exists(), rel)

    def test_the_owner_skill_headings_are_all_mapped(self):
        prov = (SKILL / "PROVENANCE.md").read_text(encoding="utf-8")
        table = prov.split("### Where each heading went", 1)[1].split("\n### ", 1)[0]
        rows = {r.split("|")[1].strip() for r in table.splitlines() if r.startswith("| ") and "---" not in r}
        for file, heads in inventory()["headings"].items():
            for h in heads:
                self.assertIn(h, rows, (file, h))

    def test_every_source_text_file_is_mapped(self):
        for owner, here in MAPPED_TEXT.items():
            self.assertIn(owner, [f["path"] for f in inventory()["files"]])
            self.assertTrue((SKILL / here).is_file(), here)
            self.assertIn(f"`{owner}`", (SKILL / "PROVENANCE.md").read_text(encoding="utf-8")
                          .replace("`references/brand-rules.md`, `SKILL.md`", "`references/brand-rules.md` `SKILL.md`"))

    def test_provenance_records_the_merge(self):
        prov = (SKILL / "PROVENANCE.md").read_text(encoding="utf-8")
        for needed in ("frq-4-pptx-agent v1.0, 2026-10-09", "kit owner", "commentAuthors", "MSIP",
                       "layouts whose rights are unclear" if False else "rights are unclear",
                       "wide ATM aircraft photo", "strip_pptx_metadata.py"):
            self.assertIn(needed, prov)
        for n in range(11, 19):
            self.assertIn(f"C{n}", prov)
        rules_ = (SKILL / "references/brand-rules.md").read_text(encoding="utf-8")
        for n in range(11, 19):
            self.assertRegex(rules_, rf"\| C{n} \|")
        self.assertEqual(rules_.count("**to confirm with GCM**"), 3)    # C11, C12, C13
        self.assertNotRegex(prov + rules_ + self.skill, r"[\w.+-]+@[\w-]+\.\w+", "no e-mail addresses")

    def test_no_new_frq_skill_name(self):
        self.assertEqual([s for s in packs.available_skills(KIT) if s.startswith("frq-")], ["frq-brandbook"])
        for p in (KIT / packs.SKILLS_REL).rglob("*"):
            if p.is_file() and "frq-brandbook" not in p.parts:
                try:
                    text = p.read_text(encoding="utf-8")
                except UnicodeDecodeError:
                    continue
                self.assertNotIn("frq-4-pptx", text, p)


class TestPlacement(unittest.TestCase):
    def test_setup_places_no_binary_brand_asset_and_the_skill_still_works(self):
        """The brand skill's binaries stay in .ai-sdlc/kit (design §8.7); the placed scripts
        find them there, and remove takes everything back."""
        with tempfile.TemporaryDirectory() as tmp:
            root = helpers.make_repo(Path(tmp).resolve() / "repo", {"README.md": "team\n"})
            before = helpers.snapshot(root)
            copy = helpers.copy_kit(root / "ai-sdlc-kit")
            kit = root / paths.KIT_REL
            helpers.cli(root, copy, "setup", "--protect-only")
            code, out = helpers.cli(root, kit, "setup", "--name", "Ana", "--roles", "po", "--lang", "en")
            self.assertEqual(code, 0, out)
            placed = [p for p in (root / PLACED).rglob("*") if p.is_file()]
            for p in placed:
                try:
                    p.read_bytes().decode("utf-8")
                except UnicodeDecodeError:
                    self.fail(f"binary file placed: {p.relative_to(root)}")
            self.assertLessEqual(sum(p.stat().st_size for p in placed), BUDGET_PLACED)
            for rel in ("assets/templates/frq-template-slim-core.pptx", "assets/templates/frq-master.pptx",
                        "assets/keyvisual/keyvisual-atm-aircraft.jpeg"):
                self.assertFalse((root / PLACED / rel).exists(), rel)
                self.assertEqual((kit / packs.SKILLS_REL / "frq-brandbook" / rel).read_bytes(),
                                 (SKILL / rel).read_bytes(), rel)
            for rel in ("assets/logo/logo-frequentis-wordmark-blue.svg", "scripts/check_brand.py",
                        "brand-tokens.json"):
                self.assertEqual((root / PLACED / rel).read_bytes(), (SKILL / rel).read_bytes(), rel)
            self.assertIn("name: ai-sdlc-frq-brandbook",
                          (root / PLACED / "SKILL.md").read_text(encoding="utf-8"))
            code, out = helpers.cli(root, kit, "check")
            self.assertNotIn(PLACED, out)
            self.assertNotIn("missing:", out)
            # the placed checker finds its tokens and assets and runs from the placed folder
            r = subprocess.run([sys.executable, "-I", str(root / PLACED / "scripts/check_brand.py"),
                                str(GOLDEN), "--json"], capture_output=True, text=True, check=False,
                               cwd=root)
            self.assertEqual(r.returncode, 1, r.stderr)
            if importlib.util.find_spec("pptx"):        # the placed builder finds the slim template
                outline = root / "outline.md"
                outline.write_text("# Towers\n\n## Controllers see more\n- One position\n", encoding="utf-8")
                deck = root / "deck.pptx"
                r = subprocess.run([sys.executable, str(root / PLACED / "scripts/new_deck.py"), str(outline),
                                    str(deck), "--classification", "Frequentis General"],
                                   capture_output=True, text=True, check=False, cwd=root)
                self.assertEqual(r.returncode, 0, r.stderr)
                self.assertEqual(count_layouts(deck), 25)
                outline.unlink()
                deck.unlink()
            code, out = helpers.cli(root, kit, "remove", "--yes")
            self.assertEqual(code, 0, out)
            self.assertFalse((root / PLACED).exists())
            self.assertEqual(helpers.snapshot(root), before)


if __name__ == "__main__":
    unittest.main()
