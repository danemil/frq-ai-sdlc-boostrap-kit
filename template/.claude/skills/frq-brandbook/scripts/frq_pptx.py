#!/usr/bin/env python3
"""Build Frequentis decks from a JSON spec on the company templates.

    python3 frq_pptx.py layouts [--template slim|full]
    python3 frq_pptx.py build SPEC.json OUT.pptx --classification "<the class the person gave>" [--template slim|full] [--year 2026]
    python3 frq_pptx.py footer IN.pptx OUT.pptx --classification "<class>" [--title …] [--presenter …] [--year …]
    python3 frq_pptx.py audit DECK.pptx [--json]        # runs check_brand.py, the one brand check
    python3 frq_pptx.py render DECK.pptx [OUTDIR]       # PNGs under .ai-sdlc/tmp/<name>/ only

From the kit owner's frq-4-pptx-agent v1.0, merged into this skill (design 2026-10-09 §8.5),
under the kit's rules:

- --classification is required: one of "Frequentis Public", "Frequentis General",
  "Frequentis Confidential", or, when the person could not be asked, the literal
  "Frequentis [classification to be set]". Ask the person; never guess it. build() and
  set_footer() take it by keyword and check it too.
- It never overwrites its input and refuses an output that already exists.
- Template: the slim template (25 layouts, about 1.2 MB per deck) by default; the full
  master (44 layouts, about 9 MB per deck) only when a slide needs a layout the slim one
  lacks, and it says so. --template slim|full forces one.
- render writes only under .ai-sdlc/tmp/ (the LibreOffice profile too), never to /tmp.
- The palette comes from ../brand-tokens.json; the templates and pictures are found through
  brand_assets.py (the placed skill first, then .ai-sdlc/kit), so this works from the placed
  folder although the binaries stay in the kit copy.

Needs python-pptx (and lxml, which it brings; Pillow only for image blocks). Run it with the
personal venv (~/.ai-sdlc/venv/bin/python) after the person agreed to install it, exactly as
ai-sdlc-doc-powerpoint describes. Run it as a command; do not import it from the repo.
"""
from __future__ import annotations

import sys

sys.dont_write_bytecode = True                         # no __pycache__ in the placed skill

import argparse
import datetime as dt
import importlib.util
import json
import shutil
import subprocess
from pathlib import Path

HERE = Path(__file__).resolve().parent


def _sibling(name):
    spec = importlib.util.spec_from_file_location(f"frq_{name}", HERE / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


brand_assets = _sibling("brand_assets")

try:                                                   # python-pptx only from the consented venv
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt
    PPTX_ERROR = None
except ImportError as exc:                             # a plain message later, never a traceback
    PPTX_ERROR = exc

NO_PPTX = ("frq_pptx: python-pptx is not installed for this Python. Ask the person before "
           "installing anything; with their yes, install it into the personal venv "
           "~/.ai-sdlc/venv as ai-sdlc-doc-powerpoint (section 1) describes, then run this "
           "script with ~/.ai-sdlc/venv/bin/python.")


def need_pptx():
    if PPTX_ERROR is not None:
        raise ImportError(NO_PPTX)


# --- brand tokens (one palette: ../brand-tokens.json) -------------------------------------
TOKENS = json.loads(brand_assets.find("brand-tokens.json").read_text(encoding="utf-8"))


def _hex(name):
    for group in ("primary", "accent"):
        for c in TOKENS["colours"][group]:
            if c["name"] == name:
                return c["hex"].lstrip("#").upper()
    raise KeyError(name)


BLUE, LIGHT_BLUE, BLACK = _hex("Blue"), _hex("Light blue"), _hex("Black")
DARK_GREY, GREY, WARM_GREY = _hex("Cool grey"), _hex("Mid grey"), _hex("Warm grey")
GREEN, ORANGE, RED, WHITE = _hex("Green"), _hex("Orange"), _hex("Red"), _hex("White")
SERIES = [h.lstrip("#").upper() for h in TOKENS["charts"]["series"]]
SERIES_RECESSED = SERIES[2:]
TRACK = TOKENS["charts"]["kpi_track"].lstrip("#").upper()
CLASSES = tuple(TOKENS["classification"])
PLACEHOLDER = TOKENS["classification_placeholder"]
ALLOWED = CLASSES + (PLACEHOLDER,)
TEMPLATES = {"slim": TOKENS["pptx"]["template"], "full": TOKENS["pptx"]["template_full"]}
AREA = TOKENS["pptx"]["content_area_in"]
ARIAL = TOKENS["fonts"]["office"]["family"]


def checked_classification(value):
    """The class the person gave, or the placeholder; anything else is refused."""
    if value not in ALLOWED:
        raise ValueError(f"classification {value!r} is not one of: " + ", ".join(ALLOWED)
                         + ". Ask the person; never guess it.")
    return value


def template_path(which="slim") -> Path:
    """The slim template or the full master, from the placed skill or .ai-sdlc/kit."""
    return brand_assets.find(TEMPLATES[which])


def rgb(hex_):
    return RGBColor.from_string(hex_)


def emu_in(v) -> float:
    return round((v or 0) / 914400, 2)


# --- layouts and the template choice ------------------------------------------------------
def layout_names(path) -> list:
    need_pptx()
    return [lay.name.strip() for lay in Presentation(str(path)).slide_layouts]


def choose_template(spec, template=None):
    """(path, 'slim' or 'full', [layouts only the full master has]) for a spec."""
    if template not in (None, "slim", "full"):
        raise ValueError(f"--template must be slim or full, not {template!r}")
    wanted = [str(s.get("layout", "")).strip() for s in spec.get("slides", [])]
    slim = {n.lower() for n in layout_names(template_path("slim"))}
    missing = sorted({n for n in wanted if n.lower() not in slim})
    if template == "slim" and missing:
        raise ValueError("the slim template does not have: " + ", ".join(missing)
                         + ". Use --template full (a deck on the full master is about 9 MB).")
    which = template or ("full" if missing else "slim")
    return template_path(which), which, (missing if which == "full" else [])


def find_layout(prs, name):
    want = name.strip().lower()
    for lay in prs.slide_layouts:
        if lay.name.strip().lower() == want:
            return lay
    names = ", ".join(repr(l.name.strip()) for l in prs.slide_layouts)
    raise ValueError(f"unknown layout {name!r}. Available: {names}")


def cmd_layouts(args):
    need_pptx()
    slim = {n.lower() for n in layout_names(template_path("slim"))}
    prs = Presentation(str(template_path(args.template or "full")))
    for i, lay in enumerate(prs.slide_layouts, 1):
        where = "slim and full" if lay.name.strip().lower() in slim else "full only"
        print(f"{i:2d}. {lay.name.strip()}  ({where})")
        for ph in lay.placeholders:
            f = ph.placeholder_format
            print(f"      idx {f.idx:<3} {str(f.type).split()[0]:<10} "
                  f"@({emu_in(ph.left)}, {emu_in(ph.top)}) {emu_in(ph.width)}x{emu_in(ph.height)} in  {ph.name}")
    return 0


# --- text helpers --------------------------------------------------------------------------
def _items(body):
    """Body spec as (text, level) tuples. A nested list is one level deeper."""
    out = []

    def walk(seq, lvl):
        for it in seq:
            if isinstance(it, list):
                walk(it, lvl + 1)
            elif isinstance(it, dict):
                out.append((str(it["text"]), int(it.get("level", lvl))))
            else:
                out.append((str(it), lvl))
    walk(body if isinstance(body, list) else [body], 0)
    return out


def fill_text(tf, body):
    items = _items(body)
    tf.text = ""
    for n, (text, lvl) in enumerate(items):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.text = text
        p.level = lvl


def style_run(run, size=None, color=None, bold=False):
    f = run.font
    f.name = ARIAL
    f.bold = bold
    f.italic = False
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = rgb(color)


# --- placeholders --------------------------------------------------------------------------
SUBTITLE_IDX = {2, 14}


def classify_placeholders(slide):
    title, subtitle, content = None, None, []
    for ph in slide.placeholders:
        f = ph.placeholder_format
        if f.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            title = ph
        elif f.idx in SUBTITLE_IDX and subtitle is None:
            subtitle = ph
        elif f.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
            content.append(ph)
    content.sort(key=lambda p: (p.left or 0, p.top or 0))
    return title, subtitle, content


def take_box(ph):
    box = (ph.left, ph.top, ph.width, ph.height)
    ph._element.getparent().remove(ph._element)
    return box


def remove_empty_placeholders(slide):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            continue
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


# --- tables --------------------------------------------------------------------------------
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"


def _cell_border(cell, side, hex_=None, w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"L": "a:lnL", "R": "a:lnR", "T": "a:lnT", "B": "a:lnB"}[side]
    for old in tcPr.findall(qn(tag)):
        tcPr.remove(old)
    ln = etree.SubElement(tcPr, qn(tag), w=str(w if hex_ else 0))
    if hex_:
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=hex_)
    else:
        etree.SubElement(ln, qn("a:noFill"))
    for e in [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill"))]:
        tcPr.remove(e)                                 # schema order: borders before fills
        tcPr.append(e)


def add_table(slide, box, spec):
    header, rows = spec["header"], spec.get("rows", [])
    left, top, width, _ = box
    row_h = Emu(int(Pt(22)))
    gf = slide.shapes.add_table(len(rows) + 1, len(header), left, top, width, Emu(row_h * (len(rows) + 1)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE_TABLE
    tbl.first_row, tbl.horz_banding = True, False
    widths = spec.get("col_widths")
    if widths:
        tot = sum(widths)
        for c, w in enumerate(widths):
            tbl.columns[c].width = Emu(int(width * w / tot))
    size = spec.get("font_size", 12)
    for r in range(len(rows) + 1):
        for c in range(len(header)):
            cell = tbl.cell(r, c)
            text = header[c] if r == 0 else (rows[r - 1][c] if c < len(rows[r - 1]) else "")
            cell.text = str(text)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(BLUE if r == 0 else WHITE)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    style_run(run, size, WHITE if r == 0 else BLACK)
            for side in "LRT":
                _cell_border(cell, side)
            _cell_border(cell, "B", GREY if r else None)
    return gf


# --- charts --------------------------------------------------------------------------------
def _chart_types():
    return {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "stacked": XL_CHART_TYPE.COLUMN_STACKED, "line": XL_CHART_TYPE.LINE,
            "doughnut": XL_CHART_TYPE.DOUGHNUT, "pie": XL_CHART_TYPE.PIE}


def _series_colours(series_specs):
    recessed = iter(SERIES_RECESSED * 4)
    out = []
    for s in series_specs:
        h = s.get("highlight")
        out.append(s.get("color") or (BLUE if h in (True, 1, "primary") else
                                      LIGHT_BLUE if h in (2, "secondary") else next(recessed)))
    if not any(s.get("highlight") or s.get("color") for s in series_specs):
        out[0] = BLUE                                  # always lead with blue
    return out


def add_chart(slide, box, spec):
    ctype = spec.get("type", "column")
    left, top, width, height = box
    if ctype == "kpi":                                 # one value on a light grey track (C12)
        v = float(spec["value"])
        cd = CategoryChartData()
        cd.categories = ["value", "rest"]
        cd.add_series("kpi", (v, max(0.0, 100 - v)))
        side = min(width, height)
        gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, left, top, side, side, cd)
        ch = gf.chart
        ch.has_legend = False
        ch.has_title = False
        plot = ch.plots[0]
        for i, col in enumerate([spec.get("color", BLUE), TRACK]):
            pt = plot.series[0].points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(col)
            pt.format.line.fill.background()
        hole = plot._element.find(qn("c:holeSize"))
        if hole is not None:
            hole.set("val", "75")
        tb = slide.shapes.add_textbox(left, top + side // 2 - Pt(14), side, Pt(28))
        tb.text_frame.text = spec.get("label", f"{v:g}%")
        p = tb.text_frame.paragraphs[0]
        p.alignment = 2                                # a number in a donut: the one centred text
        style_run(p.runs[0], 20, spec.get("color", BLUE))
        return gf
    cd = CategoryChartData()
    cd.categories = spec["categories"]
    for s in spec["series"]:
        cd.add_series(s["name"], s["values"])
    gf = slide.shapes.add_chart(_chart_types()[ctype], left, top, width, height, cd)
    ch = gf.chart
    ch.font.name, ch.font.size = ARIAL, Pt(10)
    ch.font.color.rgb = rgb(BLACK)
    plot = ch.plots[0]
    if ctype in ("doughnut", "pie"):
        cols = spec.get("point_colors") or (SERIES * 3)
        for i, pt in enumerate(plot.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(cols[i % len(cols)])
            pt.format.line.fill.background()
    else:
        for ser, col in zip(plot.series, _series_colours(spec["series"])):
            if ctype == "line":
                ser.format.line.color.rgb = rgb(col)
                ser.format.line.width = Pt(2)
                ser.smooth = False
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = rgb(col)
                ser.format.line.fill.background()
        try:
            va = ch.value_axis
            va.has_major_gridlines = False             # no gridlines (check rule chart.gridlines)
            va.visible = not spec.get("hide_value_axis", True)
            ca = ch.category_axis
            ca.format.line.color.rgb = rgb(GREY)
            ca.has_major_gridlines = False
            if ctype == "bar":
                ca.reverse_order = True
        except (ValueError, AttributeError):
            pass
        if hasattr(plot, "gap_width"):
            plot.gap_width = 80
    if spec.get("labels", True):
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.name, dl.font.size = ARIAL, Pt(10)
        dl.font.color.rgb = rgb(BLACK)
        if spec.get("number_format"):
            dl.number_format = spec["number_format"]
            dl.number_format_is_linked = False
    multi = len(spec["series"]) > 1 or ctype in ("doughnut", "pie")
    ch.has_legend = multi
    if multi:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.name, ch.legend.font.size = ARIAL, Pt(10)
    ch.has_title = False
    return gf


# --- pictures and placeholders for humans ---------------------------------------------------
def image_path(path):
    """A key visual by its kit path ('.ai-sdlc/kit/…/frq-brandbook/assets/keyvisual/…') or skill
    path ('assets/keyvisual/…'), found in the placed skill or the kit copy; else a file."""
    rel = str(path).replace("\\", "/")
    prefix = brand_assets.KIT_SKILL_REL + "/"
    if rel.startswith(prefix):
        rel = rel[len(prefix):]
    if rel.startswith("assets/"):
        return brand_assets.find(rel)
    return Path(path).expanduser()


def add_image(slide, box, spec):
    """A picture cropped to fill the box (the template's picture-in-a-shape look)."""
    left, top, width, height = box
    path = image_path(spec["path"])
    try:
        from PIL import Image                          # Pillow: only for image blocks
    except ImportError:
        raise ImportError("frq_pptx: image blocks need Pillow; ask the person, then install it into "
                          "~/.ai-sdlc/venv as ai-sdlc-doc-powerpoint describes") from None
    pic = slide.shapes.add_picture(str(path), left, top, width, height)
    with Image.open(path) as im:
        iw, ih = im.size
    box_r, img_r = width / height, iw / ih
    if img_r > box_r:
        c = (1 - box_r / img_r) / 2
        pic.crop_left = pic.crop_right = c
    else:
        c = (1 - img_r / box_r) / 2
        pic.crop_top = pic.crop_bottom = c
    pic._element.find(".//" + qn("p:cNvPr")).set("descr", spec.get("alt", Path(path).stem.replace("-", " ")))
    return pic


def flat(shape):
    """No outline, no shadow: clear effects and point the theme effect reference at 'none'."""
    shape.line.fill.background()
    shape.shadow.inherit = False
    ref = shape._element.find(".//" + qn("a:effectRef"))
    if ref is not None:
        ref.set("idx", "0")
    return shape


def add_human_placeholder(slide, box, label):
    left, top, width, height = box
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(WARM_GREY)
    flat(shp)
    shp.text_frame.text = label
    style_run(shp.text_frame.paragraphs[0].runs[0], 12, BLACK)
    return shp


def add_icon_placeholder(slide, left, top, size, descriptor):
    sq = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, size, size)
    sq.fill.solid()
    sq.fill.fore_color.rgb = rgb(BLUE)
    flat(sq)
    tb = slide.shapes.add_textbox(left, top + size + Pt(4), size * 2, Pt(16))
    tb.text_frame.text = descriptor
    style_run(tb.text_frame.paragraphs[0].runs[0], 10, BLACK)


def place_block(slide, box, block, notes):
    """block: a list or str = bullets; a dict with table, chart, image, photo_request or icons."""
    if isinstance(block, dict):
        if "table" in block:
            add_table(slide, box, block["table"])
        elif "chart" in block:
            add_chart(slide, box, block["chart"])
        elif "image" in block:
            add_image(slide, box, block["image"])
        elif "photo_request" in block:
            add_human_placeholder(slide, box, f"Photograph: {block['photo_request']}")
            notes.append("Request real photo or screenshot from Frequentis Photo Stock (GCM): "
                         + block["photo_request"])
        elif "icons" in block:
            left, top, width, _ = box
            size = Emu(int(Pt(36)))
            step = width // max(len(block["icons"]), 1)
            for i, d in enumerate(block["icons"]):
                add_icon_placeholder(slide, left + step * i, top, size, d)
                notes.append(f"Insert icon from Frequentis Icon Stock: {d}")
        elif "bullets" in block:
            return block["bullets"]
        else:
            raise ValueError(f"unknown block keys: {list(block)}")
        return None
    return block


# --- build ---------------------------------------------------------------------------------
# Boxes (inches: left, top, width, height) for layouts without a content placeholder.
FREE_BOXES = {
    "small side bar 1": {"content": (2.36, 0.83, 7.17, 4.24), "bar": (0.16, 0.83, 1.85, 4.2)},
    "small side bar 2": {"content": (2.36, 0.30, 7.17, 4.77), "bar": (0.16, 2.30, 1.85, 2.7)},
    "medium side bar 1": {"content": (3.74, 0.83, 5.79, 4.24), "bar": (0.16, 0.83, 3.2, 4.2)},
    "medium side bar 2": {"content": (3.74, 0.30, 5.79, 4.77), "bar": (0.16, 1.55, 3.2, 3.5)},
    "wide side bar": {"content": (6.70, 0.30, 2.83, 4.77), "bar": (0.47, 0.83, 5.7, 4.2)},
    "reference": {"content": (0.47, 1.10, 9.06, 3.0)},
}
DEFAULT_BOX = (AREA["left"], AREA["top"], AREA["right"] - AREA["left"], AREA["bottom"] - AREA["top"])
DEFAULT_BOX_SUB = (AREA["left"], AREA["top_with_subheadline"], AREA["right"] - AREA["left"],
                   AREA["bottom"] - AREA["top_with_subheadline"])


def inches(box):
    return tuple(Emu(int(v * 914400)) for v in box)


def build_slide(prs, sl):
    lay = find_layout(prs, sl["layout"])
    slide = prs.slides.add_slide(lay)
    title, subtitle, content = classify_placeholders(slide)
    notes = []
    if sl.get("title") is not None and title is not None:
        title.text_frame.text = sl["title"]
    if sl.get("subtitle") is not None and subtitle is not None:
        subtitle.text_frame.text = sl["subtitle"]
    for idx, text in (sl.get("ph") or {}).items():
        fill_text(slide.placeholders[int(idx)].text_frame, text)
        content = [c for c in content if c.placeholder_format.idx != int(idx)]
    blocks = sl.get("content")
    if blocks:
        if not isinstance(blocks, list) or not isinstance(blocks[0], (list, dict)):
            blocks = [blocks]                          # a single column of bullets
        if content:
            for ph, block in zip(content, blocks):
                bullets = place_block(slide, (ph.left, ph.top, ph.width, ph.height), block, notes)
                if bullets is not None:
                    fill_text(ph.text_frame, bullets)
                else:
                    take_box(ph)
        else:
            fb = FREE_BOXES.get(lay.name.strip().lower(), {})
            box = inches(fb.get("content") or (DEFAULT_BOX_SUB if subtitle else DEFAULT_BOX))
            n, gap = len(blocks), Emu(int(0.12 * 914400))
            w = (box[2] - gap * (n - 1)) // n
            for i, block in enumerate(blocks):
                b = (box[0] + (w + gap) * i, box[1], w, box[3])
                bullets = place_block(slide, b, block, notes)
                if bullets is not None:
                    tb = slide.shapes.add_textbox(*b)
                    tb.text_frame.word_wrap = True
                    fill_text(tb.text_frame, bullets)
                    for p in tb.text_frame.paragraphs:
                        for r in p.runs:
                            style_run(r, 16, BLACK)
    if sl.get("sidebar_text"):
        fb = FREE_BOXES.get(lay.name.strip().lower())
        if not fb or "bar" not in fb:
            raise ValueError(f"sidebar_text needs a side-bar layout, not {lay.name.strip()!r}")
        tb = slide.shapes.add_textbox(*inches(fb["bar"]))
        tb.text_frame.word_wrap = True
        fill_text(tb.text_frame, sl["sidebar_text"])
        for p in tb.text_frame.paragraphs:
            for r in p.runs:
                style_run(r, 14, WHITE)
    remove_empty_placeholders(slide)
    all_notes = ([sl["notes"]] if sl.get("notes") else []) + notes
    if all_notes:
        slide.notes_slide.notes_text_frame.text = "\n".join(all_notes)
    return slide


def build(spec, out, *, classification, year=None, template=None) -> dict:
    """Build a deck from `spec` into the new file `out`. Returns {"out", "template", "full_only"}."""
    classification = checked_classification(classification)  # before anything is read or written
    out = Path(out)
    if out.exists():
        raise FileExistsError(f"{out} exists; choose another name (it is never overwritten)")
    given = spec.get("classification")
    if given not in (None, classification):
        raise ValueError(f"the spec says classification {given!r} but the person gave {classification!r}")
    need_pptx()
    path, which, full_only = choose_template(spec, template)
    prs = Presentation(str(path))
    for sl in spec["slides"]:
        build_slide(prs, sl)
    first_title = next((s.get("title") for s in spec["slides"] if s.get("title")), "")
    set_footer(prs, classification=classification, year=year or spec.get("year"),
               title=spec.get("title") or first_title, presenter=spec.get("presenter") or "")
    prs.save(str(out))
    return {"out": out, "template": which, "full_only": full_only}


def cmd_build(args):
    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    result = build(spec, args.out, classification=args.classification, year=args.year,
                   template=args.template)
    print(f"wrote {result['out']} on the {result['template']} template")
    if result["full_only"]:
        print("used the full master for: " + ", ".join(result["full_only"]))
    if args.classification == PLACEHOLDER:
        print(f"frq_pptx: the footer says '{PLACEHOLDER}': the person sets the class before the deck "
              "is shared", file=sys.stderr)
    return 0


# --- footer --------------------------------------------------------------------------------
def _master_footer_boxes(prs):
    """{role: shape} of the master footer text boxes (title, presenter, class, copyright)."""
    boxes = []

    def walk(shapes):
        for s in shapes:
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(s.shapes)
            elif s.name in ("FRQ_Filename", "FRQ_Classification", "FRQ_Copyright") and s.has_text_frame:
                boxes.append(s)
    walk(prs.slide_masters[0].shapes)
    roles = {}
    for s in boxes:
        if s.name == "FRQ_Filename":
            roles["title"] = s
        elif s.name == "FRQ_Copyright":
            roles["copyright"] = s
    cls = sorted((s for s in boxes if s.name == "FRQ_Classification"), key=lambda s: s.left)
    if len(cls) == 2:                                  # the left one holds the presenter
        roles["presenter"], roles["class"] = cls
    elif cls:
        roles["class"] = cls[0]
    return roles


def _set_box_text(shape, text):
    tf = shape.text_frame
    p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.text = text
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def set_footer(prs, *, classification, year, title=None, presenter=None):
    """The master footer, text only: title | presenter | class | © Frequentis AG <year> (C11)."""
    classification = checked_classification(classification)
    year = year or dt.date.today().year
    roles = _master_footer_boxes(prs)
    if "class" not in roles:
        raise ValueError("the slide master has no classification text to set (not the Frequentis template?)")
    _set_box_text(roles["class"], classification)
    if "copyright" in roles:
        _set_box_text(roles["copyright"], f"© Frequentis AG {year}")
    if title is not None and "title" in roles:
        _set_box_text(roles["title"], title)
    if presenter is not None and "presenter" in roles:
        _set_box_text(roles["presenter"], presenter)
    return roles


def cmd_footer(args):
    src, out = Path(args.inp), Path(args.out)
    if out.resolve() == src.resolve() or out.exists():
        raise FileExistsError(f"{out} exists or is the input; write to a new file (it is never overwritten)")
    checked_classification(args.classification)
    need_pptx()
    prs = Presentation(str(src))
    roles = set_footer(prs, classification=args.classification, year=args.year, title=args.title,
                       presenter=args.presenter)
    prs.save(str(out))
    print("Footer: " + " | ".join(r.text_frame.text for r in
                                  (roles.get(k) for k in ("title", "presenter", "class", "copyright")) if r))
    return 0


# --- audit: the one check ------------------------------------------------------------------
def cmd_audit(args):
    """Runs check_brand.py (stdlib, the one brand check) and passes its result on."""
    check_brand = _sibling("check_brand")
    argv = [args.deck] + (["--json"] if args.json else []) + (["--year", str(args.year)] if args.year else [])
    return check_brand.main(argv)


# --- render: only under .ai-sdlc/tmp/ -------------------------------------------------------
def work_root(start=None) -> Path:
    """The repo root: the nearest folder up from here that holds .ai-sdlc or .git."""
    here = Path(start or Path.cwd()).resolve()
    for folder in (here, *here.parents):
        if (folder / ".ai-sdlc").is_dir() or (folder / ".git").exists():
            return folder
    return here


def cmd_render(args):
    deck = Path(args.deck).resolve()
    tmp = work_root() / ".ai-sdlc" / "tmp"
    out = Path(args.outdir).resolve() if args.outdir else tmp / deck.stem
    if out != tmp and tmp not in out.parents:
        raise ValueError(f"render writes only under {tmp} (hidden from git), not {out}")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice or not shutil.which("pdftoppm"):
        print("frq_pptx: render needs LibreOffice (soffice) and poppler (pdftoppm); without them, ask the "
              "person to look at the deck in PowerPoint", file=sys.stderr)
        return 2
    out.mkdir(parents=True, exist_ok=True)
    profile = tmp / "lo-profile"
    subprocess.run([soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless",
                    "--convert-to", "pdf", "--outdir", str(out), str(deck)],
                   check=True, capture_output=True)
    pdf = out / (deck.stem + ".pdf")
    subprocess.run(["pdftoppm", "-r", str(args.dpi), "-png", str(pdf), str(out / "slide")], check=True)
    print("\n".join(str(p) for p in sorted(out.glob("slide-*.png"))))
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Build Frequentis decks from a JSON spec.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("layouts", help="list the layouts (slim and full, or full only)")
    a.add_argument("--template", choices=["slim", "full"])
    a.set_defaults(fn=cmd_layouts)
    a = sub.add_parser("build", help="build a deck from a JSON spec into a new file")
    a.add_argument("spec")
    a.add_argument("out")
    a.add_argument("--classification", choices=ALLOWED, required=True,
                   help="ask the person; never guess it. If you cannot ask: " + repr(PLACEHOLDER))
    a.add_argument("--template", choices=["slim", "full"], help="default: slim, full only when a layout needs it")
    a.add_argument("--year", type=int)
    a.set_defaults(fn=cmd_build)
    a = sub.add_parser("footer", help="set the master footer, into a new file")
    a.add_argument("inp")
    a.add_argument("out")
    a.add_argument("--classification", choices=ALLOWED, required=True)
    a.add_argument("--title")
    a.add_argument("--presenter")
    a.add_argument("--year", type=int)
    a.set_defaults(fn=cmd_footer)
    a = sub.add_parser("audit", help="run check_brand.py, the one brand check")
    a.add_argument("deck")
    a.add_argument("--json", action="store_true")
    a.add_argument("--year", type=int)
    a.set_defaults(fn=cmd_audit)
    a = sub.add_parser("render", help="slides to PNG under .ai-sdlc/tmp/ (LibreOffice and poppler)")
    a.add_argument("deck")
    a.add_argument("outdir", nargs="?")
    a.add_argument("--dpi", type=int, default=80)
    a.set_defaults(fn=cmd_render)
    args = ap.parse_args(argv)
    try:
        return args.fn(args)
    except ImportError as exc:
        print(str(exc) if str(exc).startswith("frq_pptx") else NO_PPTX, file=sys.stderr)
    except (ValueError, KeyError, OSError, brand_assets.AssetMissing) as exc:
        print(f"frq_pptx: {exc}", file=sys.stderr)
    except subprocess.CalledProcessError as exc:
        print(f"frq_pptx: {exc.cmd[0]} failed (exit {exc.returncode})", file=sys.stderr)
    return 2


if __name__ == "__main__":
    sys.exit(main())
