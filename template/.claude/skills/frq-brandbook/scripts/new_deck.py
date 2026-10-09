#!/usr/bin/env python3
"""Build a new on-brand deck from a short Markdown outline and the bundled template.

    python new_deck.py outline.md out.pptx --classification "Frequentis General" [--author "Ana Pop"]

Needs python-pptx (run it with ~/.ai-sdlc/venv/bin/python when python3 has no pptx; install
only with the person's consent, as ai-sdlc-doc-powerpoint describes). It never overwrites
a file and never changes the template.

Outline format:

    # Deck title                      -> "Standard TITLE" (the next plain line is the sub-title)
    Ana Pop, 9 October 2026

    ## Slide headline                 -> "Headline + field", or "Sub-headline + field"
    > Optional sub-headline              when a "> " line follows
    - Bullet
      - Second-level bullet (two spaces)

A "Closing Slide" is added at the end (never a "Thank you" slide). Slides are filled by
layout name and placeholder index only: no font, size or colour is set, so the template's
theme applies. The master footer gets the classification, the year and the deck title.
"""
from __future__ import annotations

import argparse
import datetime
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parent / "assets/templates/frq-template-slim-core.pptx"
CLASSES = ("Frequentis Public", "Frequentis General", "Frequentis Confidential")


def parse(text):
    """(title, subtitle, [(headline, subheadline, [(bullet, level)])])."""
    title, subtitle, slides = "", "", []
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("## "):
            slides.append([line[3:].strip(), "", []])
        elif line.startswith("# "):
            title = line[2:].strip()
        elif slides and line.startswith("> "):
            slides[-1][1] = line[2:].strip()
        elif slides and re.match(r"^\s*[-*] ", line):
            indent = len(line) - len(line.lstrip())
            slides[-1][2].append((re.sub(r"^\s*[-*] ", "", line), min(indent // 2, 4)))
        elif not slides and title and not subtitle:
            subtitle = line.strip()
    return title, subtitle, [tuple(s) for s in slides]


def placeholder(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    raise KeyError(f"placeholder idx {idx} not on layout '{slide.slide_layout.name}'")


def fill(frame, bullets):
    frame.text = bullets[0][0] if bullets else ""
    if bullets:
        frame.paragraphs[0].level = bullets[0][1]
    for text, level in bullets[1:]:
        p = frame.add_paragraph()
        p.text, p.level = text, level


def master_runs(shapes):
    for shape in shapes:
        if shape.shape_type == 6:                      # a group: look inside
            yield from master_runs(shape.shapes)
        elif shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                yield from p.runs


def set_footer(prs, classification, year, title, author):
    """The master footer, text only (its formatting stays): class, year, title, presenter."""
    for run in master_runs(prs.slide_masters[0].shapes):
        if run.text in CLASSES:
            run.text = classification
        elif re.fullmatch(r"© Frequentis AG \d{4}", run.text):
            run.text = f"© Frequentis AG {year}"
        elif run.text == "Presentation title":
            run.text = title
        elif run.text == "<by Presenter>":
            run.text = f"by {author}" if author else ""


def build(outline, out, classification, year, author=None, template=TEMPLATE):
    from pptx import Presentation                       # python-pptx, from the consented venv

    title, subtitle, slides = parse(Path(outline).read_text(encoding="utf-8"))
    if not title:
        raise ValueError("the outline needs a '# Deck title' line")
    prs = Presentation(str(template))
    layouts = {l.name: l for l in prs.slide_layouts}
    s = prs.slides.add_slide(layouts["Standard TITLE"])
    s.shapes.title.text = title
    placeholder(s, 2).text = subtitle or (author or "")
    for headline, sub, bullets in slides:
        if sub:
            s = prs.slides.add_slide(layouts["Sub-headline + field"])
            placeholder(s, 14).text = sub
            body = placeholder(s, 16)
        else:
            s = prs.slides.add_slide(layouts["Headline + field"])
            body = placeholder(s, 15)
        s.shapes.title.text = headline
        fill(body.text_frame, bullets)
    prs.slides.add_slide(layouts["Closing Slide"])
    set_footer(prs, classification, year, title, author)
    prs.save(str(out))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="New on-brand deck from an outline and the bundled template.")
    ap.add_argument("outline")
    ap.add_argument("out")
    ap.add_argument("--classification", choices=CLASSES, default="Frequentis General")
    ap.add_argument("--year", type=int, default=datetime.date.today().year)
    ap.add_argument("--author")
    args = ap.parse_args(argv)
    if Path(args.out).exists():
        print(f"new_deck: {args.out} exists; choose another name (it is never overwritten)", file=sys.stderr)
        return 2
    try:
        build(args.outline, args.out, args.classification, args.year, args.author)
    except ImportError:
        print("new_deck: python-pptx is missing; see ai-sdlc-doc-powerpoint, section 1", file=sys.stderr)
        return 2
    except (ValueError, KeyError, OSError) as exc:
        print(f"new_deck: {exc}", file=sys.stderr)
        return 2
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
